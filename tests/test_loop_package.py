import json
import importlib
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


TEST_DIR = Path(__file__).resolve().parent
PROJECT_DIR = TEST_DIR.parent
sys.path.insert(0, str(PROJECT_DIR))


PACKAGE_DIR = PROJECT_DIR / "loop_packages" / "data_cleaning_file_organization"


class TestDataCleaningLoopPackage(unittest.TestCase):
    def test_manifest_declares_complete_loop_package_contract(self):
        from loop_packages.data_cleaning_file_organization.validators import validate_loop_package

        result = validate_loop_package(PACKAGE_DIR)

        self.assertEqual(result["status"], "pass", result)
        manifest = result["manifest"]
        self.assertEqual(manifest["schema_version"], "agentplatform.loop_package.v1")
        self.assertEqual(manifest["name"], "data_cleaning_file_organization")
        self.assertEqual(manifest["skill_interface"]["skill_name"], "data_cleaning_file_organization")
        self.assertEqual(manifest["loop_policy"]["policy_module"], "skill_policies.data_cleaning_file_organization")
        self.assertEqual(manifest["workset_schema"], "workset_schema.json")
        self.assertEqual(manifest["trace_contract"], "trace_contract.json")
        self.assertIn("validate_workset", manifest["validators"])
        self.assertIn("validate_trace_contract", manifest["validators"])
        self.assertIn("validate_run_artifacts", manifest["validators"])
        self.assertIn("fixtures/workset_minimal.json", manifest["fixtures"])
        self.assertIn("prepare_file_organization_run", manifest["runtime_tools"])
        self.assertEqual(
            manifest["runtime_registrar"],
            "_register_data_cleaning_file_organization_tools",
        )
        self.assertIn("blocked_is_not_success", manifest["hard_rules"])
        self.assertIn("needs_confirmation_stops_loop", manifest["hard_rules"])
        self.assertIn("ToolRegistry", manifest["derived_sources"])

    def test_minimal_fixture_matches_workset_schema(self):
        from loop_packages.data_cleaning_file_organization.validators import (
            load_json,
            validate_workset,
        )

        schema = load_json(PACKAGE_DIR / "workset_schema.json")
        workset = load_json(PACKAGE_DIR / "fixtures" / "workset_minimal.json")

        result = validate_workset(workset, schema)

        self.assertEqual(result["status"], "pass", result)
        self.assertEqual(workset["schema_version"], "data_cleaning.workset.v1")
        self.assertEqual(workset["task_type"], "file_organization")
        self.assertTrue(workset["inputs"]["source_files"])
        self.assertIn("structured_outputs_written", workset["done_when"])
        self.assertTrue(workset["permission_boundary"]["readonly_prepare"])
        self.assertFalse(workset["permission_boundary"]["external_write_allowed"])

    def test_trace_contract_accepts_loop_trace_shape(self):
        from loop_packages.data_cleaning_file_organization.validators import (
            load_json,
            validate_trace_contract,
        )

        trace_contract = load_json(PACKAGE_DIR / "trace_contract.json")
        trace = {
            "trace_id": "trace-test",
            "agent_name": "project_manager_agent",
            "goal": "请整理文件并归档 C:\\tmp\\采购公告.docx",
            "status": "completed",
            "metadata": {
                "active_skill": "data_cleaning_file_organization",
                "exposed_tools": ["prepare_file_organization_run"],
            },
            "rounds": [
                {
                    "action": "prepare_file_organization_run",
                    "action_input": {"file_paths": ["C:\\tmp\\采购公告.docx"]},
                    "status": "success",
                    "observation_evaluation": {
                        "schema_version": "loop.observation_evaluation.v1",
                        "status": "success",
                        "needs_confirmation": False,
                    },
                }
            ],
        }

        result = validate_trace_contract(trace, trace_contract)

        self.assertEqual(result["status"], "pass", result)

    def test_runtime_can_discover_data_cleaning_loop_package(self):
        from loop_packages import discover_loop_packages, get_loop_package

        packages = discover_loop_packages()
        package = get_loop_package("data_cleaning_file_organization")

        self.assertIn("data_cleaning_file_organization", packages)
        self.assertEqual(package.name, "data_cleaning_file_organization")
        self.assertEqual(package.manifest["schema_version"], "agentplatform.loop_package.v1")
        self.assertEqual(package.workset_schema["workset_schema_version"], "data_cleaning.workset.v1")
        self.assertEqual(package.trace_contract["schema_version"], "agentplatform.trace_contract.v1")

    def test_all_runtime_skills_have_loop_packages(self):
        from loop_packages import discover_loop_packages
        import main

        packages = discover_loop_packages()

        self.assertEqual(set(packages), set(main.SKILL_TOOL_MAP))
        for skill_name, expected_tools in main.SKILL_TOOL_MAP.items():
            with self.subTest(skill_name=skill_name):
                package = packages[skill_name]
                self.assertEqual(package.skill_name, skill_name)
                self.assertEqual(package.expected_tools, expected_tools)
                self.assertTrue(package.manifest["route_keywords"])

    def test_all_loop_packages_have_executable_validators_and_fixtures(self):
        from loop_packages import discover_loop_packages

        packages = discover_loop_packages()

        for skill_name, package in packages.items():
            with self.subTest(skill_name=skill_name):
                self.assertTrue(package.manifest["validators"])
                self.assertTrue(package.manifest["fixtures"])
                validators = importlib.import_module(f"loop_packages.{skill_name}.validators")
                manifest_result = validators.validate_loop_package(package.path)
                self.assertEqual(manifest_result["status"], "pass", manifest_result)

                workset_schema = validators.load_json(package.path / package.manifest["workset_schema"])
                fixture_path = package.path / package.manifest["fixtures"][0]
                workset = validators.load_json(fixture_path)
                workset_result = validators.validate_workset(workset, workset_schema)
                self.assertEqual(workset_result["status"], "pass", workset_result)

    def test_validate_all_loop_packages_reports_runtime_alignment(self):
        from loop_packages import validate_all_loop_packages

        result = validate_all_loop_packages()

        self.assertEqual(result["status"], "pass", result)
        self.assertEqual(result["package_count"], 4)
        self.assertEqual(result["errors"], [])

    def test_governance_validator_exposes_loop_package_target(self):
        import governance.validate as governance_validate

        errors, warnings, findings = governance_validate.validate_loop_packages(verbose=False)

        self.assertEqual(errors, 0, findings)
        self.assertEqual(warnings, 0, findings)

    def test_governance_cli_accepts_loop_packages_command(self):
        result = subprocess.run(
            [sys.executable, "-X", "utf8", "-B", "governance/validate.py", "loop-packages"],
            cwd=PROJECT_DIR,
            text=True,
            capture_output=True,
            check=False,
        )

        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn("loop package", result.stdout.lower())
        self.assertIn("0 errors, 0 warnings", result.stdout)

    def test_package_route_keywords_drive_runtime_skill_selection(self):
        from loop_packages import route_skill_from_packages
        import main

        self.assertEqual(route_skill_from_packages("crm_submit_gate"), "cloudcc_crm")
        self.assertEqual(route_skill_from_packages("bid_notice_intake"), "opportunity_management")
        self.assertEqual(route_skill_from_packages("project_status_rollup"), "project_management")
        self.assertEqual(main._route_skill("crm_submit_gate"), "cloudcc_crm")
        self.assertEqual(main._route_skill("bid_notice_intake"), "opportunity_management")
        self.assertEqual(main._route_skill("project_status_rollup"), "project_management")

    def test_runtime_registry_matches_loop_package_tool_contract(self):
        from loop_packages import get_loop_package
        import main

        package = get_loop_package("data_cleaning_file_organization")
        reg = main._build_registry_for_skill(package.skill_name)

        self.assertEqual(package.expected_tools, package.manifest["runtime_tools"])
        self.assertEqual(reg.list_tools(), package.expected_tools)

    def test_main_skill_tool_map_is_derived_from_loop_packages(self):
        from loop_packages import build_skill_tool_map
        import main

        main_source = (PROJECT_DIR / "main.py").read_text(encoding="utf-8")

        self.assertEqual(main.SKILL_TOOL_MAP, build_skill_tool_map())
        self.assertNotIn("SKILL_TOOL_MAP = {", main_source)
        self.assertIn("build_skill_tool_map", main_source)

    def test_main_registrar_map_is_derived_from_loop_packages(self):
        from loop_packages import build_skill_registrar_map
        import main

        main_source = (PROJECT_DIR / "main.py").read_text(encoding="utf-8")
        expected = build_skill_registrar_map(main)

        self.assertEqual(main.SKILL_REGISTRARS, expected)
        self.assertNotIn("SKILL_REGISTRARS = {", main_source)
        self.assertIn("build_skill_registrar_map", main_source)

    def test_prepare_run_artifacts_match_package_contract(self):
        from docx import Document
        from loop_packages.data_cleaning_file_organization.validators import validate_run_artifacts
        from tools.data_cleaning_tools import DataCleaningTools

        with tempfile.TemporaryDirectory() as td:
            source_dir = os.path.join(td, "source")
            workspace_dir = os.path.join(td, "workspace")
            os.makedirs(source_dir)

            doc_path = os.path.join(source_dir, "purchase_notice.docx")
            doc = Document()
            doc.add_paragraph("Test customer")
            doc.add_paragraph("Project name:")
            doc.add_paragraph("Printer procurement project")
            doc.save(doc_path)

            result = DataCleaningTools(workspace_dir=workspace_dir).prepare_file_organization_run([doc_path])
            validation = validate_run_artifacts(result)

            self.assertEqual(validation["status"], "pass", validation)
            self.assertTrue(os.path.exists(result["artifacts"]["input_manifest"]))
            self.assertTrue(os.path.exists(result["artifacts"]["review_queue"]))
            self.assertTrue(os.path.exists(result["artifacts"]["planned_archive_actions"]))
            self.assertTrue(os.path.exists(result["artifacts"]["run_report"]))
            self.assertTrue(os.path.exists(result["artifacts"]["trace"]))

    def test_real_main_run_trace_replays_against_package_contract(self):
        from docx import Document
        from loop_packages import get_loop_package
        from loop_packages.data_cleaning_file_organization.validators import (
            load_json,
            validate_trace_contract,
        )
        import main

        with tempfile.TemporaryDirectory() as td:
            source_dir = os.path.join(td, "source")
            workspace_dir = os.path.join(td, "workspace")
            trace_dir = os.path.join(td, "logs")
            os.makedirs(source_dir)

            doc_path = os.path.join(source_dir, "purchase_notice.docx")
            doc = Document()
            doc.add_paragraph("Test customer")
            doc.add_paragraph("Project name:")
            doc.add_paragraph("Printer procurement project")
            doc.save(doc_path)

            result = main.run(
                f"请整理文件并归档 {doc_path}",
                planner_mode="rule",
                data_workspace_dir=workspace_dir,
                trace_dir=trace_dir,
            )

            package = get_loop_package("data_cleaning_file_organization")
            trace_path = os.path.join(trace_dir, f"project_manager_{result['trace_id']}.json")
            saved_trace = load_json(Path(trace_path))
            contract_result = validate_trace_contract(saved_trace, package.trace_contract)

            self.assertEqual(result["metadata"]["active_skill"], package.skill_name)
            self.assertEqual(contract_result["status"], "pass", contract_result)


if __name__ == "__main__":
    unittest.main()
