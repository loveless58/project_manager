"""DocxExecutor 真实样本测试。"""

import os
import tempfile
import unittest

from docx import Document

from integrations.executors import DocxExecutor


class TestDocxExecutor(unittest.TestCase):

    def setUp(self):
        self.ex = DocxExecutor()
        self.path = None

    def tearDown(self):
        if self.path and os.path.exists(self.path):
            os.unlink(self.path)

    def _create_test_docx(self, paragraphs=None, table_rows=None, author=None, title=None):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            self.path = f.name
        doc = Document()
        for p in (paragraphs or []):
            doc.add_paragraph(p)
        if table_rows:
            table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
            for i, row in enumerate(table_rows):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = str(cell)
        if author:
            doc.core_properties.author = author
        if title:
            doc.core_properties.title = title
        doc.save(self.path)

    # ---- can_handle ----

    def test_can_handle_docx_extension(self):
        self.assertTrue(self.ex.can_handle("/tmp/a.docx"))

    def test_can_handle_docx_uppercase(self):
        self.assertTrue(self.ex.can_handle("/tmp/a.DOCX"))

    def test_cannot_handle_xlsx(self):
        self.assertFalse(self.ex.can_handle("/tmp/a.xlsx"))

    def test_cannot_handle_pdf(self):
        self.assertFalse(self.ex.can_handle("/tmp/a.pdf"))

    def test_cannot_handle_url(self):
        self.assertFalse(self.ex.can_handle("https://example.com/a.docx"))

    # ---- extract success ----

    def test_extract_basic(self):
        self._create_test_docx(
            paragraphs=["段落1", "段落2", ""],  # 空段落过滤
            author="测试作者",
            title="测试标题",
        )
        result = self.ex.extract(self.path)
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["implementation_status"], "implemented")
        self.assertEqual(len(result["raw_data"]["paragraphs"]), 2)
        self.assertEqual(result["raw_data"]["metadata"]["author"], "测试作者")
        self.assertEqual(result["raw_data"]["metadata"]["title"], "测试标题")

    def test_extract_with_tables(self):
        self._create_test_docx(
            paragraphs=["表格测试"],
            table_rows=[["A", "B"], ["1", "2"]],
        )
        result = self.ex.extract(self.path)
        self.assertEqual(len(result["raw_data"]["tables"]), 1)
        self.assertEqual(result["raw_data"]["tables"][0], [["A", "B"], ["1", "2"]])

    def test_extract_empty_table_rows_filtered(self):
        self._create_test_docx(paragraphs=["无表文档"])
        result = self.ex.extract(self.path)
        self.assertEqual(result["raw_data"]["tables"], [])

    def test_extract_raw_text_includes_paragraphs_and_tables(self):
        self._create_test_docx(
            paragraphs=["段落A", "段落B"],
            table_rows=[["col1", "col2"], ["v1", "v2"]],
        )
        result = self.ex.extract(self.path)
        text = result["raw_data"]["raw_text"]
        self.assertIn("段落A", text)
        self.assertIn("段落B", text)
        self.assertIn("col1", text)
        self.assertIn("v1", text)

    # ---- extract error ----

    def test_extract_file_not_found(self):
        result = self.ex.extract("/nonexistent.docx")
        self.assertEqual(result["status"], "blocked")
        self.assertIn("not found", result["error"].lower())

    def test_extract_corrupted_file(self):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            f.write(b"not a real docx")
            self.path = f.name
        result = self.ex.extract(self.path)
        self.assertEqual(result["status"], "blocked")
        self.assertIn("parse error", result["error"].lower())

    # ---- get_supported_types ----

    def test_get_supported_types(self):
        self.assertEqual(self.ex.get_supported_types(), [".docx"])
