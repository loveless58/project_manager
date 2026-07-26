from __future__ import annotations

import hashlib
import io
import json
import math
import re
from pathlib import Path
from unittest.mock import patch

import pytest

from platform_core.models import BusinessContextEvidence
from platform_core.storage_bindings import StorageBinding, StorageBindingRegistry


CANDIDATE_ID = f"C-{1:03d}"
CONTRACT_CODE = f"CT-{1:03d}"


class StaticRetrieval:
    def __init__(self, evidence: BusinessContextEvidence) -> None:
        self.evidence = evidence
        self.calls = 0

    def find_business_candidates(self, query):
        self.calls += 1
        return self.evidence


class SchemaValidInterpreter:
    name = "fake"
    model = "fake-model"
    schema_version = "candidate_document_interpretation.v1"
    prompt_version = "document_interpretation.v1"
    policy_version = "document_interpretation_policy.v1"

    def __init__(self) -> None:
        self.calls = 0

    def complete_json(self, request):
        self.calls += 1
        return {
            "schema_version": self.schema_version,
            "status": "success",
            "document_type": "contract",
            "fields": {"contract_code": CONTRACT_CODE},
            "relations": [{"relation_type": "contract_project", "target_candidate_id": CANDIDATE_ID}],
            "evidence": [{"kind": "business_context", "candidate_id": CANDIDATE_ID, "field": "contract_code"}],
            "confidence": 0.94,
            "interpreter": self.name,
            "model": self.model,
            "prompt_version": self.prompt_version,
            "policy_version": self.policy_version,
        }


class CountingInterpretation:
    def __init__(self) -> None:
        self.calls = 0

    def interpret(self, evidence_pack):
        self.calls += 1
        raise AssertionError("ambiguous binding must stop before interpretation")


def _binding(binding_id: str, root: Path, roles: tuple[str, ...], *, readable: bool = True,
             writable: bool = False, enabled: bool = True) -> StorageBinding:
    return StorageBinding(binding_id, "local", "node", f"business://{binding_id}/", root,
                          roles, readable, writable, enabled)


def _matched_context() -> BusinessContextEvidence:
    return BusinessContextEvidence(
        "matched",
        ({"id": "C-001", "document_type": "contract", "parties": {},
          "facts": {"contract_code": CONTRACT_CODE}},),
        ({"kind": "business_context", "candidate_id": CANDIDATE_ID, "field": "contract_code"},),
        (),
        ({"code": "BUSINESS_CONTEXT.CANDIDATES_FOUND", "count": 1},),
    )


def test_preparation_builds_review_only_interpretation_and_intent_artifacts(tmp_path: Path) -> None:
    from integrations.document_store import DocumentStoreRouter, LocalDocumentStore
    from services.archive_targets import ArchiveTargetResolver
    from services.document_interpretation import DocumentInterpretationService
    from tools.data_cleaning_tools import DataCleaningTools

    source_root, target_root, workspace = tmp_path / "source", tmp_path / "archive", tmp_path / "runtime"
    source_root.mkdir()
    target_root.mkdir()
    source = source_root / "contract.md"
    original = f"合同\n合同编号：{CONTRACT_CODE}\n项目名称：合成项目001"
    source.write_text(original, encoding="utf-8")
    registry = StorageBindingRegistry([
        _binding("source", source_root, ("source",)),
        _binding("archive", target_root, ("archive_target",), readable=False, writable=True),
    ])
    router = DocumentStoreRouter(registry, {"source": LocalDocumentStore(source_root)})
    retrieval = StaticRetrieval(_matched_context())
    interpreter = SchemaValidInterpreter()
    interpretation = DocumentInterpretationService(retrieval, interpreter)
    tools = DataCleaningTools(
        workspace_dir=str(workspace), storage_binding_registry=registry,
        document_store_router=router, retrieval_service=retrieval,
        interpretation_service=interpretation,
        archive_target_resolver=ArchiveTargetResolver(registry),
    )

    with patch.object(router, "open_read", wraps=router.open_read) as routed_read:
        result = tools.prepare_file_organization_run([str(source)])

    assert routed_read.call_count == 1
    assert re.fullmatch(r"run_[0-9a-f]{32}", result["run_id"])
    interpretation_item = result["candidate_interpretations"][0]
    intent_item = result["archive_intents"][0]
    action_item = result["archive_actions"][0]
    assert interpretation_item["run_id"] == intent_item["run_id"] == action_item["run_id"] == result["run_id"]
    assert interpretation_item["source_ref"] == intent_item["source_ref"] == action_item["source_ref"]
    assert interpretation_item["content_hash"] == intent_item["content_hash"]
    assert interpretation_item["parse_artifact_ref"].endswith(interpretation_item["content_hash"][:24])
    assert result["candidate_interpretations"][0]["business_relation"]["candidate_contract_id"] == CANDIDATE_ID
    assert result["archive_intents"][0]["schema_version"] == "archive_intent.v1"
    assert result["archive_intents"][0]["destination_status"] == "unresolved"
    assert result["archive_actions"][0]["status"] == "needs_review"
    assert result["archive_actions"][0]["confirmed"] is False
    assert result["archive_actions"][0]["archive_decision"]["human_review_required"] is True
    assert result["archive_actions"][0]["archive_decision"]["target_dir"] is None
    assert result["archive_actions"][0]["archive_decision"]["target_path"] is None
    assert Path(result["artifacts"]["candidate_interpretations"]).is_file()
    assert Path(result["artifacts"]["archive_intents"]).is_file()
    assert source.read_text(encoding="utf-8") == original
    assert not list(tmp_path.rglob("archive_result.json"))
    assert not list(tmp_path.rglob("project_ledger.json"))
    assert retrieval.calls == 1
    assert interpreter.calls == 1


def test_prepare_invoice_under_project_path_keeps_strict_artifact_project_free(
    tmp_path: Path,
) -> None:
    from docx import Document
    from integrations.document_store import DocumentStoreRouter, LocalDocumentStore
    from services.archive_targets import ArchiveTargetResolver
    from services.document_interpretation import DocumentInterpretationService
    from tools.data_cleaning_tools import DataCleaningTools

    class InvoiceInterpreter:
        name = "fake"
        model = "fake-model"
        schema_version = "candidate_document_interpretation.v1"
        prompt_version = "document_interpretation.v1"
        policy_version = "document_interpretation_policy.v1"

        def complete_json(self, request):
            return {
                "schema_version": self.schema_version,
                "status": "success",
                "document_type": "invoice",
                "fields": {},
                "relations": [
                    {
                        "relation_type": "invoice_contract",
                        "target_candidate_id": CANDIDATE_ID,
                    }
                ],
                "evidence": [
                    {
                        "kind": "business_context",
                        "candidate_id": CANDIDATE_ID,
                        "field": "contract_code",
                    }
                ],
                "confidence": 0.94,
                "interpreter": self.name,
                "model": self.model,
                "prompt_version": self.prompt_version,
                "policy_version": self.policy_version,
            }

    source_root = tmp_path / "source"
    source = source_root / "项目执行" / "合成项目001" / "原始文件" / "invoice.docx"
    source.parent.mkdir(parents=True)
    document = Document()
    document.add_paragraph("Invoice")
    document.add_paragraph("Invoice Number: SYN-INVOICE-STRICT-001")
    document.add_paragraph("项目名称：合成项目999")
    document.save(source)
    archive_root = tmp_path / "archive"
    archive_root.mkdir()
    registry = StorageBindingRegistry([
        _binding("source", source_root, ("source",)),
        _binding(
            "archive", archive_root, ("archive_target",),
            readable=False, writable=True,
        ),
    ])
    router = DocumentStoreRouter(
        registry, {"source": LocalDocumentStore(source_root)}
    )
    retrieval = StaticRetrieval(_matched_context())
    interpretation = DocumentInterpretationService(
        retrieval, InvoiceInterpreter()
    )
    tools = DataCleaningTools(
        workspace_dir=str(tmp_path / "runtime"),
        storage_binding_registry=registry,
        document_store_router=router,
        retrieval_service=retrieval,
        interpretation_service=interpretation,
        archive_target_resolver=ArchiveTargetResolver(registry),
    )

    result = tools.prepare_file_organization_run([str(source)])
    artifact = json.loads(Path(result["structured_outputs"][0]).read_text(encoding="utf-8"))

    assert artifact["document_type"] == "invoice"
    assert "project_name" not in artifact["candidate_fields"]
    assert "lifecycle_stage" not in artifact["candidate_fields"]
    assert artifact["classification"]["business_domain"] == "finance"
    assert artifact["classification"]["requires_review"] is True
    assert result["candidate_interpretations"][0]["status"] == "needs_review"
    assert result["archive_actions"][0]["status"] == "needs_review"
    assert result["archive_actions"][0]["confirmed"] is False


def test_runtime_workspace_must_stay_outside_storage_bindings(tmp_path: Path) -> None:
    from services.archive_targets import ArchiveTargetResolver
    from tools.data_cleaning_tools import DataCleaningTools

    binding_root = tmp_path / "sync-root"
    workspace = binding_root / "runtime"
    registry = StorageBindingRegistry([
        _binding("source", binding_root, ("source",)),
    ])

    with pytest.raises(ValueError, match="runtime workspace must be outside storage bindings"):
        DataCleaningTools(
            workspace_dir=str(workspace),
            storage_binding_registry=registry,
            archive_target_resolver=ArchiveTargetResolver(registry),
        )


def test_nested_binding_ambiguity_stops_before_parse_retrieval_and_interpretation(tmp_path: Path) -> None:
    from services.archive_targets import ArchiveTargetResolver
    from tools.data_cleaning_tools import DataCleaningTools

    shared, incoming = tmp_path / "shared", tmp_path / "shared" / "incoming"
    incoming.mkdir(parents=True)
    source = incoming / "contract.md"
    source.write_text("合同", encoding="utf-8")
    registry = StorageBindingRegistry([
        _binding("shared", shared, ("source",)),
        _binding("incoming", incoming, ("source",)),
    ])
    retrieval = StaticRetrieval(_matched_context())
    interpretation = CountingInterpretation()
    tools = DataCleaningTools(
        workspace_dir=str(tmp_path / "runtime"), storage_binding_registry=registry,
        retrieval_service=retrieval, interpretation_service=interpretation,
        archive_target_resolver=ArchiveTargetResolver(registry),
    )

    with patch.object(tools, "extract_document", wraps=tools.extract_document) as native_parse:
        result = tools.prepare_file_organization_run([str(source)])

    assert "STORAGE_BINDING.AMBIGUOUS" in json.dumps(result, ensure_ascii=False)
    assert native_parse.call_count == retrieval.calls == interpretation.calls == 0
    assert result["candidate_interpretations"] == []
    assert result["archive_actions"][0]["status"] == "needs_review"
    assert not list(tmp_path.rglob("archive_result.json"))


class ReplacingRouter:
    def __init__(self, delegate, source: Path) -> None:
        self.delegate = delegate
        self.source = source
        self.calls = 0

    def open_read(self, ref):
        self.calls += 1
        with self.delegate.open_read(ref) as stream:
            snapshot = stream.read()
        self.source.write_text("replacement", encoding="utf-8")
        return io.BytesIO(snapshot)


class MaliciousInterpretation:
    def __init__(self, payload: dict) -> None:
        self.payload = payload

    def interpret(self, evidence_pack):
        return dict(self.payload)


def _configured_tools(tmp_path: Path, *, retrieval=None, interpretation=None, router_wrapper=None):
    from integrations.document_store import DocumentStoreRouter, LocalDocumentStore
    from services.archive_targets import ArchiveTargetResolver
    from tools.data_cleaning_tools import DataCleaningTools

    source_root = tmp_path / "source"
    source_root.mkdir()
    source = source_root / "contract.md"
    source.write_text(f"合同\n合同编号：{CONTRACT_CODE}", encoding="utf-8")
    registry = StorageBindingRegistry([
        _binding("source", source_root, ("source",)),
        _binding("archive", tmp_path / "archive", ("archive_target",), readable=False, writable=True),
    ])
    router = DocumentStoreRouter(registry, {"source": LocalDocumentStore(source_root)})
    if router_wrapper is not None:
        router = router_wrapper(router, source)
    retrieval = retrieval or StaticRetrieval(_matched_context())
    tools = DataCleaningTools(
        workspace_dir=str(tmp_path / "runtime"), storage_binding_registry=registry,
        document_store_router=router, retrieval_service=retrieval,
        interpretation_service=interpretation,
        archive_target_resolver=ArchiveTargetResolver(registry),
    )
    return tools, source, retrieval, router


@pytest.mark.skipif(not hasattr(Path, "is_symlink"), reason="symlink unavailable")
def test_source_symlink_is_rejected_before_parse_retrieval_and_interpretation(tmp_path: Path) -> None:
    target = tmp_path / "target.md"
    target.write_text("合同", encoding="utf-8")
    source_root = tmp_path / "source"
    source_root.mkdir()
    source = source_root / "contract.md"
    try:
        source.symlink_to(target)
    except OSError as exc:
        pytest.skip(f"symlink unavailable: {exc}")
    retrieval = StaticRetrieval(_matched_context())
    interpretation = CountingInterpretation()
    from integrations.document_store import DocumentStoreRouter, LocalDocumentStore
    from services.archive_targets import ArchiveTargetResolver
    from tools.data_cleaning_tools import DataCleaningTools
    registry = StorageBindingRegistry([_binding("source", source_root, ("source",))])
    tools = DataCleaningTools(
        workspace_dir=str(tmp_path / "runtime"), storage_binding_registry=registry,
        document_store_router=DocumentStoreRouter(registry, {"source": LocalDocumentStore(source_root)}),
        retrieval_service=retrieval, interpretation_service=interpretation,
        archive_target_resolver=ArchiveTargetResolver(registry),
    )

    with patch.object(tools, "extract_document", wraps=tools.extract_document) as native_parse:
        result = tools.prepare_file_organization_run([str(source)])

    assert "DOCUMENT_SOURCE.SYMLINK_REJECTED" in json.dumps(result, ensure_ascii=False)
    assert native_parse.call_count == retrieval.calls == interpretation.calls == 0
    assert not result["structured_outputs"]


def test_source_replacement_during_single_router_snapshot_fails_closed(tmp_path: Path) -> None:
    interpretation = CountingInterpretation()
    tools, source, retrieval, router = _configured_tools(
        tmp_path, interpretation=interpretation, router_wrapper=ReplacingRouter,
    )

    with patch.object(tools, "extract_document", wraps=tools.extract_document) as native_parse:
        result = tools.prepare_file_organization_run([str(source)])

    assert "DOCUMENT_SOURCE.REPLACED" in json.dumps(result, ensure_ascii=False)
    assert router.calls == 1
    assert native_parse.call_count == retrieval.calls == interpretation.calls == 0
    assert not result["structured_outputs"]


def test_malicious_interpretation_output_is_rejected_and_not_persisted(tmp_path: Path) -> None:
    sentinel = "unexpected-field-sentinel"
    payload = {
        "schema_version": "candidate_document_interpretation.v1",
        "status": "success",
        "document_type": "contract",
        "fields": {"contract_code": CONTRACT_CODE},
        "relations": [{"relation_type": "contract_project", "target_candidate_id": CANDIDATE_ID}],
        "evidence": [{"kind": "business_context", "candidate_id": CANDIDATE_ID, "field": "contract_code"}],
        "confidence": 0.9,
        "interpreter": "fake", "model": "fake-model",
        "prompt_version": "document_interpretation.v1",
        "policy_version": "document_interpretation_policy.v1",
        "unknown": sentinel,
    }
    tools, source, _, _ = _configured_tools(
        tmp_path, interpretation=MaliciousInterpretation(payload),
    )

    result = tools.prepare_file_organization_run([str(source)])

    persisted = Path(result["artifacts"]["candidate_interpretations"]).read_text(encoding="utf-8")
    assert result["candidate_interpretations"][0]["status"] == "blocked"
    assert result["candidate_interpretations"][0]["blocked_reason"] == "DOCUMENT_INTERPRETATION.SCHEMA_INVALID"
    assert sentinel not in persisted


def test_malicious_retrieval_diagnostics_are_rejected_and_not_persisted(tmp_path: Path) -> None:
    unsafe_value = "Bearer" + " " + "-".join(("synthetic", "credential"))
    context = BusinessContextEvidence(
        "matched",
        ({"id": CANDIDATE_ID, "document_type": "contract", "parties": {},
          "facts": {"contract_code": CONTRACT_CODE}},),
        ({"kind": "business_context", "candidate_id": CANDIDATE_ID, "field": "contract_code"},),
        (),
        ({"code": "BUSINESS_CONTEXT.CANDIDATES_FOUND", "credential": unsafe_value},),
    )
    tools, source, _, _ = _configured_tools(
        tmp_path, retrieval=StaticRetrieval(context), interpretation=MaliciousInterpretation({}),
    )

    result = tools.prepare_file_organization_run([str(source)])

    all_json = "\n".join(
        item.read_text(encoding="utf-8") for item in Path(result["artifacts"]["run_dir"]).glob("*.json")
    )
    assert result["candidate_interpretations"][0]["business_context"]["status"] == "blocked"
    assert "BUSINESS_CONTEXT.SCHEMA_INVALID" in all_json
    assert unsafe_value not in all_json


def test_registered_intent_run_cannot_be_downgraded_to_legacy_execution(tmp_path: Path) -> None:
    tools, source, _, _ = _configured_tools(
        tmp_path, interpretation=MaliciousInterpretation({}),
    )
    prepared = tools.prepare_file_organization_run([str(source)])
    run_id = prepared["run_id"]
    run_dir = Path(prepared["artifacts"]["run_dir"])
    target = tmp_path / "legacy-target" / source.name
    (run_dir / "planned_archive_actions.json").write_text(json.dumps({
        "schema_version": "archive_plan.v1",
        "run_id": run_id,
        "actions": [{
            "schema_version": "archive_action.v1", "run_id": run_id,
            "status": "ready", "source_file": str(source),
            "target_path": str(target), "blockers": [],
        }],
    }), encoding="utf-8")

    result = tools.execute_archive_plan(run_id, confirmed=True)

    assert result["status"] == "blocked"
    assert result["gate"]["blockers"] == ["archive_intent_not_executable"]
    assert source.exists()
    assert not target.exists()
    assert not (run_dir / "archive_result.json").exists()


def test_interpretation_boundary_rejects_nan_deep_and_absolute_path(tmp_path: Path) -> None:
    from contracts.archive_run_artifacts import normalize_interpretation_output

    base = {
        "schema_version": "candidate_document_interpretation.v1",
        "status": "success", "document_type": "contract",
        "fields": {"contract_code": CONTRACT_CODE},
        "relations": [{"relation_type": "contract_project", "target_candidate_id": CANDIDATE_ID}],
        "evidence": [{"kind": "business_context", "candidate_id": CANDIDATE_ID, "field": "contract_code"}],
        "confidence": 0.9, "interpreter": "fake", "model": "fake-model",
        "prompt_version": "document_interpretation.v1",
        "policy_version": "document_interpretation_policy.v1",
    }
    deep = value = {}
    for _ in range(40):
        value["child"] = {}
        value = value["child"]
    variants = []
    nan_payload = dict(base)
    nan_payload["confidence"] = float("nan")
    variants.append(nan_payload)
    path_payload = dict(base)
    path_payload["fields"] = {"contract_name": str(tmp_path.resolve())}
    variants.append(path_payload)
    deep_payload = dict(base)
    deep_payload["unexpected"] = deep
    variants.append(deep_payload)

    for payload in variants:
        result = normalize_interpretation_output(
            payload, parse_artifact_ref="artifact:parsed:0123456789abcdef01234567",
        )
        assert result["status"] == "blocked"
        assert result["blocked_reason"] == "DOCUMENT_INTERPRETATION.SCHEMA_INVALID"
        assert all(not (isinstance(item, float) and not math.isfinite(item)) for item in result.values())


def test_task5_collection_schema_rejects_unknown_root_field() -> None:
    from contracts.archive_run_artifacts import (
        ArchiveRunArtifactError,
        validate_task5_collection_artifact,
    )

    run_id = "run_44444444444444444444444444444444"
    with pytest.raises(ArchiveRunArtifactError, match="root fields"):
        validate_task5_collection_artifact(
            {
                "schema_version": "candidate_interpretations.v1",
                "run_id": run_id,
                "items": [],
                "unexpected": True,
            },
            kind="candidate_interpretations",
            run_id=run_id,
        )


def test_atomic_json_publish_rejects_nan_without_overwriting_existing_file(tmp_path: Path) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    target = tmp_path / "artifact.json"
    target.write_text('{"stable":true}', encoding="utf-8")
    tools = DataCleaningTools(workspace_dir=str(tmp_path / "runtime"))

    with pytest.raises(ValueError):
        tools._save_structured_json(str(target), {"score": float("nan")})

    assert target.read_text(encoding="utf-8") == '{"stable":true}'
    assert not list(tmp_path.glob(".artifact-*"))


def test_business_candidate_id_matching_archive_binding_is_not_a_storage_selection(tmp_path: Path) -> None:
    from integrations.document_store import DocumentStoreRouter, LocalDocumentStore
    from services.archive_targets import ArchiveTargetResolver
    from services.document_interpretation import DocumentInterpretationService
    from tools.data_cleaning_tools import DataCleaningTools

    source_root = tmp_path / "source"
    source_root.mkdir()
    source = source_root / "contract.md"
    source.write_text(f"合同\n合同编号：{CONTRACT_CODE}", encoding="utf-8")
    registry = StorageBindingRegistry([
        _binding("source", source_root, ("source",)),
        _binding(CANDIDATE_ID, tmp_path / "archive", ("archive_target",), readable=False, writable=True),
    ])
    retrieval = StaticRetrieval(_matched_context())
    tools = DataCleaningTools(
        workspace_dir=str(tmp_path / "runtime"),
        storage_binding_registry=registry,
        document_store_router=DocumentStoreRouter(registry, {"source": LocalDocumentStore(source_root)}),
        retrieval_service=retrieval,
        interpretation_service=DocumentInterpretationService(retrieval, SchemaValidInterpreter()),
        archive_target_resolver=ArchiveTargetResolver(registry),
    )

    result = tools.prepare_file_organization_run([str(source)])

    intent = result["archive_intents"][0]
    assert intent["project_id"] == CANDIDATE_ID
    assert intent["destination_status"] == "unresolved"
    assert intent["candidate_target_binding_ids"] == []
    assert "ARCHIVE_TARGET.UNRESOLVED" in intent["blockers"]


def test_agent_run_without_explicit_gateway_is_capability_disabled(tmp_path: Path) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    tools = DataCleaningTools(workspace_dir=str(tmp_path / "runtime"))

    result = tools.prepare_agent_judgement_run([str(tmp_path / "not-used.md")])

    assert result["status"] == "blocked"
    assert result["blocked_reason"] == "LLM.CAPABILITY_DISABLED"


@pytest.mark.parametrize("malicious_kind", ["path", "credential", "deep", "oversized", "nan"])
def test_malicious_native_parse_is_rejected_before_extracted_artifact_persistence(
    tmp_path: Path, malicious_kind: str,
) -> None:
    tools, source, retrieval, _ = _configured_tools(
        tmp_path, interpretation=CountingInterpretation(),
    )
    unsafe_value = ""
    fields: dict[str, object] = {"contract_code": CONTRACT_CODE}
    text_value = "合同"
    if malicious_kind == "path":
        unsafe_value = str((tmp_path / "physical-source").resolve())
        fields["contract_name"] = unsafe_value
    elif malicious_kind == "credential":
        unsafe_value = "Bearer" + " " + "-".join(("synthetic", "credential"))
        fields["contract_name"] = unsafe_value
    elif malicious_kind == "deep":
        nested: dict[str, object] = {}
        cursor = nested
        for _ in range(40):
            child: dict[str, object] = {}
            cursor["child"] = child
            cursor = child
        fields["line_items"] = [nested]
        unsafe_value = "child"
    elif malicious_kind == "oversized":
        unsafe_value = "Z" * 200_000
        text_value = unsafe_value
    else:
        fields["amount"] = float("nan")
    parser_result = {
        "schema_version": "document.extract.v1",
        "document_type": "合同",
        "classification": None,
        "fields": fields,
        "extracted_text": text_value,
    }

    with patch.object(tools, "extract_document", return_value=parser_result):
        result = tools.prepare_file_organization_run([str(source)])

    run_dir = Path(result["artifacts"]["run_dir"])
    persisted = "\n".join(
        artifact.read_text(encoding="utf-8") for artifact in run_dir.glob("*.json")
    )
    assert result["status"] == "failed"
    assert result["failures"][0]["blocked_reason"] == "DOCUMENT_PARSE.SCHEMA_INVALID"
    assert result["structured_outputs"] == []
    assert not list((run_dir / "extracted").glob("*_extracted.json"))
    assert result["candidate_interpretations"] == []
    assert result["archive_intents"] == []
    assert retrieval.calls == 0
    if unsafe_value and len(unsafe_value) < 10_000:
        assert unsafe_value not in persisted
    assert "NaN" not in persisted


@pytest.mark.parametrize(
    "malformed_kind",
    [
        "list_root", "null_root", "unknown_root", "missing_version",
        "wrong_version", "unknown_field", "line_items", "nested_field",
    ],
)
def test_native_parse_contract_rejects_unrecognized_structure_before_any_downstream_effect(
    tmp_path: Path, malformed_kind: str,
) -> None:
    tools, source, retrieval, _ = _configured_tools(
        tmp_path, interpretation=CountingInterpretation(),
    )
    parser_result: dict[str, object] = {
        "schema_version": "document.extract.v1",
        "document_type": "合同",
        "classification": None,
        "fields": {"contract_code": CONTRACT_CODE},
        "extracted_text": "合同",
    }
    fields = parser_result["fields"]
    assert isinstance(fields, dict)
    parser_output: object = parser_result
    if malformed_kind == "list_root":
        parser_output = []
    elif malformed_kind == "null_root":
        parser_output = None
    elif malformed_kind == "unknown_root":
        parser_result["parser_debug"] = "not-contract-data"
    elif malformed_kind == "missing_version":
        parser_result.pop("schema_version")
    elif malformed_kind == "wrong_version":
        parser_result["schema_version"] = "document.extract.v0"
    elif malformed_kind == "unknown_field":
        fields["unsupported_native_field"] = "not-contract-data"
    elif malformed_kind == "line_items":
        fields["line_items"] = [{"item_name": "not-contract-data"}]
    else:
        fields["buyer"] = {"name": "not-contract-data"}

    with patch.object(tools, "extract_document", return_value=parser_output):
        result = tools.prepare_file_organization_run([str(source)])

    run_dir = Path(result["artifacts"]["run_dir"])
    assert result["status"] == "failed"
    assert result["failures"][0]["blocked_reason"] == "DOCUMENT_PARSE.SCHEMA_INVALID"
    assert result["structured_outputs"] == []
    assert not list((run_dir / "extracted").glob("*_extracted.json"))
    assert result["candidate_interpretations"] == []
    assert result["archive_intents"] == []
    assert retrieval.calls == 0


@pytest.mark.parametrize(
    ("candidate_ids", "binding_options", "expected_status"),
    [
        (("source",), {}, "unresolved"),
        (("SynologyDrive",), {}, "unresolved"),
        (("disabled",), {"enabled": False}, "unresolved"),
        (("readonly",), {"writable": False}, "unresolved"),
        (("wrong-role",), {"roles": ("source",)}, "unresolved"),
        (("archive",), {}, "resolved"),
        (("archive", "archive-2"), {}, "unresolved"),
    ],
)
def test_archive_target_resolution_is_explicit_unique_and_never_defaults(
    tmp_path: Path, candidate_ids: tuple[str, ...], binding_options: dict[str, object],
    expected_status: str,
) -> None:
    from platform_core.models import DocumentRef
    from services.archive_targets import ArchiveTargetResolver

    bindings = [_binding("source", tmp_path / "source", ("source",))]
    for binding_id in {"SynologyDrive", "disabled", "readonly", "wrong-role", "archive", "archive-2"}:
        options = binding_options if binding_id in candidate_ids and len(candidate_ids) == 1 else {}
        bindings.append(_binding(
            binding_id, tmp_path / binding_id, options.get("roles", ("archive_target",)),
            readable=False, writable=options.get("writable", True), enabled=options.get("enabled", True),
        ))
    resolver = ArchiveTargetResolver(StorageBindingRegistry(bindings))
    source_ref = DocumentRef("local", "contract.md", "business://source/contract.md", "source")

    result = resolver.resolve(source_ref, candidate_ids)

    assert result.destination_status == expected_status
    if expected_status == "resolved":
        assert result.candidate_target_binding_ids == ("archive",)
    else:
        assert result.resolved_binding_id == ""
