from __future__ import annotations

import hashlib
import json
from pathlib import Path
import pytest


def _hashes(paths):
    return {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths}


def test_business_judgement_usable_slice_is_read_only_and_blocks_scans(tmp_path):
    fitz = pytest.importorskip("fitz")
    from ocr.providers import DisabledOcrProvider
    from tools.data_cleaning_tools import DataCleaningTools
    source, runtime = tmp_path / "source", tmp_path / "runtime"
    source.mkdir()
    pdf = source / "invoice.pdf"
    text_pdf = fitz.open()
    text_pdf.new_page().insert_text((72, 72), "synthetic invoice " * 30)
    text_pdf.save(pdf)
    text_pdf.close()
    from docx import Document
    docx = source / "contract.docx"
    document_docx = Document()
    document_docx.add_paragraph("Synthetic contract CT-001")
    document_docx.save(docx)
    from openpyxl import Workbook
    xlsx = source / "project.xlsx"
    workbook = Workbook()
    workbook.active.append(["project", "synthetic"])
    workbook.save(xlsx)
    markdown = source / "governance.md"
    markdown.write_text("# governance\nsynthetic review input", encoding="utf-8")
    xml = source / "project.xml"
    xml.write_text("<project><name>synthetic</name></project>", encoding="utf-8")
    scan = source / "scan.pdf"
    document = fitz.open()
    document.new_page().draw_rect((72, 72, 160, 160), fill=(0.5, 0.5, 0.5))
    document.save(scan)
    document.close()
    originals = [pdf, docx, xlsx, markdown, xml, scan]
    before = _hashes(originals)
    tools = DataCleaningTools(workspace_dir=str(runtime), ocr_adapter=DisabledOcrProvider().extract)
    result = tools.prepare_file_organization_run([str(path) for path in originals])
    assert before == _hashes(originals)
    assert result["failed"] == 1
    assert result["failures"][0]["blocked_reason"] == "OCR.CAPABILITY_DISABLED"
    assert result["processed"] == 5
    assert all(action["status"] == "needs_review" for action in result["archive_actions"])
    assert not list(runtime.rglob("archive_result.json"))


class _Task8Response:
    def __init__(self, payload):
        self.payload = payload

    def raise_for_status(self):
        return None

    def json(self):
        return self.payload


class _Task8PageIndex:
    provider_version = "synthetic-pageindex-v1"

    def __init__(self, events):
        self.events = events

    def index_pdf(self, _source_path, **kwargs):
        self.events.append("pageindex")
        return {
            "status": "success",
            "engine": "pageindex",
            "doc_name": "synthetic.pdf",
            "doc_id": "synthetic-pageindex-doc",
            "operation_id": kwargs["operation_id"],
            "provider_version": self.provider_version,
            "content_hash": kwargs["expected_content_hash"],
            "structure": [],
        }


def _task8_binding(binding_id, root, roles, *, readable, writable):
    from platform_core.storage_bindings import StorageBinding

    return StorageBinding(
        binding_id, "local", "synthetic-node", f"business://{binding_id}/", root,
        roles, readable, writable,
    )


def _task8_pdf(path, text="", *, scanned=False):
    fitz = pytest.importorskip("fitz")
    document = fitz.open()
    page = document.new_page()
    if scanned:
        page.draw_rect((72, 72, 160, 160), fill=(0.5, 0.5, 0.5))
    else:
        import html

        page.insert_htmlbox(
            page.rect,
            f"<p>{html.escape(text).replace(chr(10), "<br>")}</p>",
        )
    document.save(path)
    document.close()


def _task8_llm(events):
    def transport(**kwargs):
        events.append("llm")
        request = json.loads(kwargs["json"]["messages"][1]["content"])
        document_type = request["document"]["document_type_hint"]
        evidence_field = request["business_context"]["evidence"][0]["field"]
        if document_type == "invoice":
            response_fields = {"invoice_number": "SYN-INV-001"}
            relation = "invoice_contract"
        elif document_type == "contract":
            response_fields = {"contract_code": "CT-001"}
            relation = "contract_project"
        else:
            response_fields = {"project_code": "SYN-PROJECT-001"}
            relation = "contract_project"
        response = {
            "schema_version": "candidate_document_interpretation.v1",
            "status": "success",
            "document_type": document_type,
            "fields": response_fields,
            "relations": [{"relation_type": relation, "target_candidate_id": "C-001"}],
            "evidence": [{
                "kind": "business_context",
                "candidate_id": "C-001",
                "field": evidence_field,
            }],
            "confidence": 0.91,
            "interpreter": "openai_compatible",
            "model": "synthetic-model",
            "prompt_version": "document_interpretation.v1",
            "policy_version": "document_interpretation_policy.v1",
        }
        return _Task8Response({
            "choices": [{"message": {"content": json.dumps(response)}}]
        })

    return transport


class _Task8CliInterpreter:
    name = "synthetic-cli"
    model = "synthetic-cli-model"
    schema_version = "candidate_document_interpretation.v1"
    prompt_version = "document_interpretation.v1"
    policy_version = "document_interpretation_policy.v1"

    def complete_json(self, request):
        fields = request["document"]["candidate_fields"]
        return {
            "schema_version": self.schema_version,
            "status": "success",
            "document_type": "contract",
            "fields": {"contract_code": fields["contract_code"]},
            "relations": [{
                "relation_type": "contract_project",
                "target_candidate_id": "C-001",
            }],
            "evidence": [{
                "kind": "business_context",
                "candidate_id": "C-001",
                "field": "contract_code",
            }],
            "confidence": 0.91,
            "interpreter": self.name,
            "model": self.model,
            "prompt_version": self.prompt_version,
            "policy_version": self.policy_version,
        }


def _task8_cli_module():
    import importlib.util

    root = Path(__file__).resolve().parents[2]
    spec = importlib.util.spec_from_file_location(
        "task8_prepare_business_file_run",
        root / "scripts" / "prepare_business_file_run.py",
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_business_judgement_usable_slice_exercises_real_boundaries(tmp_path, monkeypatch, capsys):
    """Run the usable slice through real parsers/writers and fake only network edges."""
    from docx import Document
    from integrations.business_context import JsonBusinessContextProvider
    from integrations.document_store import DocumentStoreRouter, LocalDocumentStore
    from integrations.llm import OpenAICompatibleInterpreter
    from integrations.pageindex import PageIndexStructureIndex
    from ocr.providers import DisabledOcrProvider
    from openpyxl import Workbook
    from platform_core.storage_bindings import StorageBindingRegistry
    from services.archive_targets import ArchiveTargetResolver
    from services.document_interpretation import DocumentInterpretationService
    from services.retrieval_service import RetrievalService
    from tools.data_cleaning_tools import DataCleaningTools

    source, target, runtime, projections = (
        tmp_path / "source",
        tmp_path / "archive",
        tmp_path / "runtime",
        tmp_path / "projections",
    )
    for directory in (source, target, runtime, projections):
        directory.mkdir()

    invoice = source / "invoice.pdf"
    _task8_pdf(
        invoice,
        "\u7535\u5b50\u53d1\u7968\n\u53d1\u7968\u53f7\u7801\uff1aSYN-INV-001\n\u5f00\u7968\u65e5\u671f\uff1a2026-07-25\n"
        "\u8d2d\u4e70\u65b9\u540d\u79f0\uff1a\u5408\u6210\u7532\u65b9\u6709\u9650\u516c\u53f8\n\u8d2d\u4e70\u65b9\u7a0e\u53f7\uff1a913100000000000001\n"
        "\u9500\u552e\u65b9\u540d\u79f0\uff1a\u5408\u6210\u4e59\u65b9\u6709\u9650\u516c\u53f8\n\u9500\u552e\u65b9\u7a0e\u53f7\uff1a913100000000000002\n"
        "\u9879\u76ee\u540d\u79f0 | \u89c4\u683c\u578b\u53f7 | \u91d1\u989d\n\u6280\u672f\u670d\u52a1 | \u6807\u51c6\u7248 | 100.00"
        + chr(65) * 120,
    )
    contract = source / "contract.docx"
    document = Document()
    document.add_paragraph("\u6280\u672f\u5f00\u53d1\u5408\u540c")
    document.add_paragraph("\u5408\u540c\u767b\u8bb0\u7f16\u53f7\uff1aCT-001")
    document.add_paragraph("\u9879\u76ee\u7f16\u53f7\uff1aSYN-PROJECT-001")
    document.save(contract)
    governance = source / "PRD-project-manager-governance.md"
    governance.write_text("# synthetic governance\n\u9879\u76ee\u7f16\u53f7\uff1aSYN-PROJECT-001", encoding="utf-8")
    project = source / "project.xlsx"
    workbook = Workbook()
    workbook.active.append(["\u9879\u76ee\u7f16\u53f7\uff1aSYN-PROJECT-001"])
    workbook.save(project)
    xml = source / "project.xml"
    xml.write_text("<project><\u9879\u76ee\u7f16\u53f7>SYN-PROJECT-001</\u9879\u76ee\u7f16\u53f7></project>", encoding="utf-8")
    context_pdf = source / "context-contract.pdf"
    _task8_pdf(context_pdf, "synthetic context")
    scan = source / "scan.pdf"
    _task8_pdf(scan, scanned=True)

    catalog_payload = json.loads(
        (Path(__file__).resolve().parents[1] / "fixtures" / "business_context_catalog.v1.json").read_text(encoding="utf-8")
    )
    catalog_payload["records"][0]["facts"]["project_code"] = "SYN-PROJECT-001"
    catalog_payload["records"][0]["documents"][0]["path"] = str(context_pdf)
    catalog_payload["records"][0]["documents"][0]["content_hash"] = hashlib.sha256(context_pdf.read_bytes()).hexdigest()
    catalog = tmp_path / "business_context_catalog.v1.json"
    catalog.write_text(json.dumps(catalog_payload, ensure_ascii=False), encoding="utf-8")
    originals = [invoice, contract, governance, project, xml, context_pdf, scan]
    before = _hashes(originals)

    registry = StorageBindingRegistry([
        _task8_binding("source", source, ("source",), readable=True, writable=False),
        _task8_binding("archive", target, ("archive_target",), readable=False, writable=True),
    ])
    assert source.resolve() != target.resolve()
    assert registry.document_ref_from_path(invoice).binding_id == "source"

    events = []

    class TrackingCatalog:
        def __init__(self, provider):
            self.provider = provider

        def search(self, query):
            events.append("retrieval")
            return self.provider.search(query)

    context = TrackingCatalog(JsonBusinessContextProvider(catalog, registry))
    pageindex = PageIndexStructureIndex(
        tmp_path / "fake-pageindex",
        client_factory=lambda _root: _Task8PageIndex(events),
    )
    retrieval = RetrievalService(context, pageindex)
    interpreter = OpenAICompatibleInterpreter(
        base_url="https://synthetic.invalid/v1",
        api_key="fake-key-for-test",
        model="synthetic-model",
        transport=_task8_llm(events),
    )
    tools = DataCleaningTools(
        workspace_dir=str(runtime),
        ocr_adapter=DisabledOcrProvider().extract,
        storage_binding_registry=registry,
        document_store_router=DocumentStoreRouter(
            registry, {"source": LocalDocumentStore(source)}
        ),
        retrieval_service=retrieval,
        interpretation_service=DocumentInterpretationService(retrieval, interpreter),
        archive_target_resolver=ArchiveTargetResolver(registry),
    )
    def forbidden_ocr(*_args, **_kwargs):
        raise AssertionError("A real OCR provider must not be invoked.")

    for method in (
        "_default_ocr_adapter",
        "_ocr_with_easyocr",
        "_ocr_with_rapidocr",
        "_ocr_with_tesseract",
        "_ocr_with_vision_macos",
    ):
        monkeypatch.setattr(DataCleaningTools, method, staticmethod(forbidden_ocr))

    result = tools.prepare_file_organization_run([
        str(invoice), str(contract), str(governance), str(project), str(xml), str(scan)
    ])
    assert before == _hashes(originals)
    assert result["processed"] == 5, result["failures"]
    assert result["failed"] == 1
    assert result["failures"][0]["blocked_reason"] == "OCR.CAPABILITY_DISABLED"
    assert all(
        item["status"] == "needs_review"
        for item in result["candidate_interpretations"]
    ), result["candidate_interpretations"]
    assert all(action["status"] == "needs_review" and not action["confirmed"] for action in result["archive_actions"])
    assert events.count("retrieval") == events.count("pageindex") == events.count("llm") == 5
    assert all(events[index - 1] == "retrieval" for index, value in enumerate(events) if value == "pageindex")
    assert not list(runtime.rglob("archive_result.json"))

    extracted = [
        json.loads(Path(path).read_text(encoding="utf-8"))
        for path in result["structured_outputs"]
    ]
    invoice_parse = next(
        item for item in extracted
        if item["source_ref"]["object_key"] == "invoice.pdf"
    )
    invoice_interpretation = next(
        item for item in result["candidate_interpretations"]
        if item["source_ref"]["object_key"] == "invoice.pdf"
    )
    assert "project_name" not in invoice_parse["candidate_fields"]
    assert invoice_interpretation["business_relation"]["candidate_contract_id"] == "C-001"
    assert all(item["id"] and item["question"] for item in result["review_queue"]["items"])

    config = tmp_path / "node-config.json"
    config.write_text(json.dumps({
        "deployment_mode": "local",
        "runtime_workspace": str(tmp_path / "cli-runtime"),
        "storage_bindings": [
            {"binding_id": "source", "provider": "local", "node_id": "synthetic-node", "logical_root": "business://source/", "physical_root": str(source), "roles": ["source"], "readable": True, "writable": False},
            {"binding_id": "archive", "provider": "local", "node_id": "synthetic-node", "logical_root": "business://archive/", "physical_root": str(target), "roles": ["archive_target"], "readable": False, "writable": True},
        ],
        "providers": {"document_store": "local", "structure_index": "disabled", "projection_writer": "filesystem", "projection_root": str(projections)},
    }), encoding="utf-8")
    cli = _task8_cli_module()
    native_exit = cli.main([
        "--config", str(config), "--context", str(catalog),
        "--source-binding", "source", "--target-binding", "archive", str(contract),
    ], interpreter_factory=_Task8CliInterpreter)
    native_summary = json.loads(capsys.readouterr().out)
    assert native_exit == 0
    assert native_summary["status"] == "needs_review"
    assert native_summary["processed"] == 1 and native_summary["failed"] == 0
    assert native_summary["source_binding"] == "source"
    assert native_summary["target_binding"] == "archive"
    assert native_summary["archive_execution"]["confirmed"] is False

    scan_exit = cli.main([
        "--config", str(config), "--context", str(catalog),
        "--source-binding", "source", str(scan),
    ], interpreter_factory=_Task8CliInterpreter)
    scan_summary = json.loads(capsys.readouterr().out)
    assert scan_exit == 2
    assert scan_summary["status"] == "blocked"
    assert scan_summary["processed"] == 0 and scan_summary["failed"] == 1
    assert scan_summary["failure_codes"] == ["OCR.CAPABILITY_DISABLED"]
    assert not list((tmp_path / "cli-runtime").rglob("archive_result.json"))
