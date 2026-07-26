from __future__ import annotations

import json
import os
from pathlib import Path
import subprocess
import pytest


class StrictSyntheticAgentHost:
    """Test host that can respond only by reading the persisted request file."""

    def __init__(self) -> None:
        self.request_hashes: list[str] = []

    def write_responses(self, request_path: Path, response_path: Path) -> None:
        requests = json.loads(request_path.read_text(encoding="utf-8"))
        assert requests["schema_version"] == "agent_judgement_requests.v1"
        assert requests["requests"]
        self.request_hashes = [item["request_hash"] for item in requests["requests"]]
        responses = [self._response(item) for item in requests["requests"]]
        response_path.write_text(
            json.dumps(
                {
                    "schema_version": "agent_judgement_responses.v1",
                    "run_id": requests["run_id"],
                    "responses": responses,
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

    @staticmethod
    def _response(request: dict) -> dict:
        interpretation_request = request["interpretation_request"]
        document_type = interpretation_request["document"]["document_type_hint"]
        evidence = interpretation_request["business_context"]["evidence"][:1]
        return {
            "schema_version": "agent_judgement_response.v1",
            "run_id": request["run_id"],
            "request_id": request["request_id"],
            "request_hash": request["request_hash"],
            "interpreter": "strict_synthetic_agent",
            "model": "deterministic_fixture_v1",
            "created_at": "2026-07-26T00:00:00Z",
            "interpretation": {
                "schema_version": "candidate_document_interpretation.v1",
                "status": "success",
                "document_type": document_type,
                "fields": {},
                "relations": [
                    {
                        "relation_type": (
                            "invoice_contract"
                            if document_type == "invoice"
                            else "contract_project"
                        ),
                        "target_candidate_id": "C-001",
                    }
                ],
                "evidence": evidence,
                "confidence": 0.94,
                "interpreter": "strict_synthetic_agent",
                "model": "deterministic_fixture_v1",
                "prompt_version": "document_interpretation.v1",
                "policy_version": "document_interpretation_policy.v1",
            },
        }


def test_local_agent_dry_run_round_trip_is_review_only(tmp_path, monkeypatch) -> None:
    """Dropping the persisted host hand-off or a review-only gate is a bug."""
    from integrations.llm import OpenAICompatibleInterpreter
    from tools.data_cleaning_tools import DataCleaningTools

    def forbidden(*_args, **_kwargs):
        raise AssertionError("configured LLM and real OCR providers are forbidden")

    monkeypatch.setattr(OpenAICompatibleInterpreter, "__init__", forbidden)
    for method in (
        "_default_ocr_adapter",
        "_ocr_with_easyocr",
        "_ocr_with_rapidocr",
        "_ocr_with_tesseract",
        "_ocr_with_vision_macos",
    ):
        monkeypatch.setattr(DataCleaningTools, method, staticmethod(forbidden))

    original_execute = DataCleaningTools.execute_archive_plan
    archive_confirmations = []

    def record_archive_gate(self, run_id, confirmed=False):
        archive_confirmations.append(confirmed)
        return original_execute(self, run_id, confirmed=confirmed)

    monkeypatch.setattr(DataCleaningTools, "execute_archive_plan", record_archive_gate)

    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / "dry-run"
    host = StrictSyntheticAgentHost()

    result = run_local_agent_dry_run(root, host=host)

    assert result["native_statuses"] == {
        "invoice": "needs_review",
        "contract": "needs_review",
        "markdown": "needs_review",
        "xlsx": "needs_review",
        "xml": "needs_review",
    }
    assert result["scan_failure"] == "OCR.CAPABILITY_DISABLED"
    assert result["invoice_candidate"] == "C-001"
    assert result["hashes_before"] == result["hashes_after"]
    assert len(host.request_hashes) == 5
    assert result["review_items_valid"]
    assert all(result["review_items_valid"])
    assert archive_confirmations == [False]
    assert result["archive_execution"]["confirmed"] is False
    assert not list((root / "runtime").rglob("archive_result.json"))

    config = json.loads(
        (root / "config" / "project-manager.local.json").read_text(encoding="utf-8")
    )
    assert [item["binding_id"] for item in config["storage_bindings"]] == [
        "dry-run-source"
    ]
    sqlite_path = Path(config["database"]["sqlite_path"]).resolve()
    assert sqlite_path.is_file()
    assert sqlite_path.is_relative_to((root / "runtime").resolve())
    assert not sqlite_path.is_relative_to((root / "source").resolve())

    run_dir = root / "runtime" / "runs" / result["run_id"]
    plan = json.loads(
        (run_dir / "planned_archive_actions.json").read_text(encoding="utf-8")
    )
    assert plan["actions"]
    assert all(action["confirmed"] is False for action in plan["actions"])


class MissingResponseHost:
    def write_responses(self, request_path: Path, _response_path: Path) -> None:
        assert json.loads(request_path.read_text(encoding="utf-8"))["requests"]


class MutatingResponseHost(StrictSyntheticAgentHost):
    def __init__(self, mutation) -> None:
        super().__init__()
        self.mutation = mutation

    def write_responses(self, request_path: Path, response_path: Path) -> None:
        super().write_responses(request_path, response_path)
        payload = json.loads(response_path.read_text(encoding="utf-8"))
        self.mutation(payload)
        response_path.write_text(
            json.dumps(payload, ensure_ascii=False), encoding="utf-8"
        )


def _forbid_external_providers(monkeypatch) -> None:
    from integrations.llm import OpenAICompatibleInterpreter
    from tools.data_cleaning_tools import DataCleaningTools

    def forbidden(*_args, **_kwargs):
        raise AssertionError("configured LLM and real OCR providers are forbidden")

    monkeypatch.setattr(OpenAICompatibleInterpreter, "__init__", forbidden)
    for method in (
        "_default_ocr_adapter",
        "_ocr_with_easyocr",
        "_ocr_with_rapidocr",
        "_ocr_with_tesseract",
        "_ocr_with_vision_macos",
    ):
        monkeypatch.setattr(DataCleaningTools, method, staticmethod(forbidden))


def test_invalid_or_missing_host_response_fails_closed(tmp_path, monkeypatch) -> None:
    """Accepting absent, hash-altered, unsafe, or unevidenced output is a bug."""
    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    cases = [
        (
            "missing",
            MissingResponseHost(),
            "LLM.CAPABILITY_DISABLED",
        ),
        (
            "altered-hash",
            MutatingResponseHost(
                lambda payload: payload["responses"][0].__setitem__(
                    "request_hash", "0" * 64
                )
            ),
            "AGENT_JUDGEMENT.RESPONSE_INVALID",
        ),
        (
            "unsafe-path",
            MutatingResponseHost(
                lambda payload: payload["responses"][0]["interpretation"][
                    "fields"
                ].__setitem__("contract_name", "C:/unsafe/source.docx")  # repo-hygiene: allow=synthetic-path
            ),
            "AGENT_JUDGEMENT.RESPONSE_INVALID",
        ),
        (
            "relation-without-evidence",
            MutatingResponseHost(
                lambda payload: payload["responses"][0]["interpretation"].__setitem__(
                    "evidence", []
                )
            ),
            "AGENT_JUDGEMENT.RESPONSE_INVALID",
        ),
    ]

    for name, host, expected_code in cases:
        root = tmp_path / name
        result = run_local_agent_dry_run(root, host=host)

        assert result["status"] == "blocked"
        assert result["error_code"] == expected_code
        assert result["hashes_before"] == result["hashes_after"]
        assert result["archive_execution"]["confirmed"] is False
        assert not list((root / "runtime").rglob("archive_result.json"))
        assert '"confirmed": true' not in json.dumps(result).lower()


def test_source_outside_explicit_root_is_rejected_before_bootstrap(
    tmp_path, monkeypatch
) -> None:
    """Resolving an explicit source outside root/source is a safety bug."""
    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    outside = tmp_path / "outside.md"
    outside.write_text("synthetic but out of scope", encoding="utf-8")
    root = tmp_path / "dry-run"

    result = run_local_agent_dry_run(
        root, host=StrictSyntheticAgentHost(), source_paths=[outside]
    )

    assert result["status"] == "blocked"
    assert result["error_code"] == "SOURCE_BINDING.OUT_OF_SCOPE"
    assert not root.exists()


def test_existing_source_requires_explicit_synthetic_reuse_flag(
    tmp_path, monkeypatch
) -> None:
    """Silently taking ownership of an existing source directory is a bug."""
    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / "dry-run"
    (root / "source").mkdir(parents=True)

    result = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.SOURCE_ALREADY_EXISTS"
    assert not (root / "runtime").exists()


def test_unrelated_existing_root_content_is_not_modified(tmp_path) -> None:
    """Bootstrapping over an unrelated directory can overwrite user data."""
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / "not-a-dry-run"
    root.mkdir()
    sentinel = root / "keep.txt"
    sentinel.write_text("do not modify", encoding="utf-8")

    result = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.ROOT_NOT_EMPTY"
    assert sentinel.read_text(encoding="utf-8") == "do not modify"
    assert not (root / "source").exists()


def test_explicit_flag_reuses_only_recorded_synthetic_source(
    tmp_path, monkeypatch
) -> None:
    """Rejecting a hash-bound synthetic source after explicit reuse is a bug."""
    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / "dry-run"
    first = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())
    second = run_local_agent_dry_run(
        root,
        host=StrictSyntheticAgentHost(),
        use_existing_synthetic_source=True,
    )

    assert first["status"] == second["status"] == "needs_review"
    assert first["hashes_before"] == second["hashes_before"]
    assert not list((root / "runtime").rglob("archive_result.json"))


def test_cli_requires_root_and_redacts_generated_paths(
    tmp_path, monkeypatch, capsys
) -> None:
    """Making root optional or printing the physical root is a CLI safety bug."""
    import pytest

    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import main

    with pytest.raises(SystemExit) as missing_root:
        main([])
    assert missing_root.value.code == 2

    root = tmp_path / "cli-dry-run"
    code = main(["--root", str(root)])
    output = capsys.readouterr().out
    payload = json.loads(output)

    assert code == 0
    assert payload["status"] == "needs_review"
    assert str(root) not in output
    assert not list((root / "runtime").rglob("archive_result.json"))


def _make_directory_junction(link: Path, target: Path) -> None:
    completed = subprocess.run(
        ["cmd.exe", "/d", "/c", "mklink", "/J", str(link), str(target)],
        capture_output=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr or completed.stdout
    assert os.path.isjunction(link)


def _tree_hashes(root: Path) -> dict[str, str]:
    import hashlib

    return {
        path.relative_to(root).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def test_root_symlink_is_rejected_before_external_target_is_touched(
    tmp_path, monkeypatch
) -> None:
    """Resolving a root symlink before rejecting it can redirect every write."""
    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    external = tmp_path / "external-root"
    external.mkdir()
    sentinel = external / "sentinel.bin"
    sentinel.write_bytes(b"do-not-touch")
    root_link = tmp_path / "dry-run-link"
    root_link.symlink_to(external, target_is_directory=True)
    before = _tree_hashes(external)

    result = run_local_agent_dry_run(
        root_link, host=StrictSyntheticAgentHost()
    )

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.PATH_REPARSE"
    assert _tree_hashes(external) == before
    assert not (external / "source").exists()


@pytest.mark.parametrize("linked_name", ["source", "runtime", "config", "archive"])
def test_fixed_child_junction_is_rejected_before_external_writes(
    tmp_path, monkeypatch, linked_name
) -> None:
    """Following a fixed-child junction can expose source or runtime state."""
    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / f"dry-run-{linked_name}"
    first = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())
    assert first["status"] == "needs_review"

    external = tmp_path / f"external-{linked_name}"
    child = root / linked_name
    if child.exists():
        child.rename(external)
    else:
        external.mkdir()
        (external / "sentinel.bin").write_bytes(b"do-not-touch")
    _make_directory_junction(child, external)
    before = _tree_hashes(external)

    result = run_local_agent_dry_run(
        root,
        host=StrictSyntheticAgentHost(),
        use_existing_synthetic_source=True,
        include_archive_binding=linked_name == "archive",
    )

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.PATH_REPARSE"
    assert _tree_hashes(external) == before
    assert not list(external.rglob("archive_result.json"))


@pytest.mark.parametrize(
    "network_root",
    [
        r"\\synthetic-server\synthetic-share\dry-run",  # repo-hygiene: allow=synthetic-path
        "//synthetic-server/synthetic-share/dry-run",  # repo-hygiene: allow=synthetic-path
        "smb://synthetic-server/synthetic-share/dry-run",
        "nfs://synthetic-server/synthetic-share/dry-run",
        "afp://synthetic-server/synthetic-share/dry-run",
    ],
)
def test_obvious_network_root_is_rejected_before_filesystem_access(
    monkeypatch, network_root
) -> None:
    """Any filesystem call before a lexical network-root block can touch remote I/O."""
    import scripts.run_local_agent_dry_run as module

    filesystem_calls = []

    def forbidden(*args, **_kwargs):
        filesystem_calls.append(args)
        raise AssertionError("network root reached filesystem access")

    monkeypatch.setattr(module.os, "lstat", forbidden)
    monkeypatch.setattr(module, "Path", forbidden)

    result = module.run_local_agent_dry_run(
        network_root, host=StrictSyntheticAgentHost()
    )

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.ROOT_NETWORK"
    assert filesystem_calls == []


def test_reuse_rejects_forged_source_and_self_authored_marker(
    tmp_path, monkeypatch
) -> None:
    """A marker inside source cannot be the trust anchor for source reuse."""
    import hashlib

    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / "dry-run"
    first = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())
    assert first["status"] == "needs_review"

    source = root / "source"
    project_code = "".join(("PRJ", "-001"))
    contract_code = "".join(("CT", "-001"))
    forged = source / "synthetic-governance.md"
    forged.write_text(
        f"# Synthetic governance\n\n项目编号：{project_code}\n合同编号：{contract_code}\n"
        "伪造内容：marker 不能为自己授权\n",
        encoding="utf-8",
    )
    marker = source / ".local-agent-dry-run.synthetic.v1.json"
    marker_payload = json.loads(marker.read_text(encoding="utf-8"))
    marker_payload["hashes"][forged.name] = hashlib.sha256(
        forged.read_bytes()
    ).hexdigest()
    marker.write_text(json.dumps(marker_payload), encoding="utf-8")
    runtime_before = _tree_hashes(root / "runtime")

    result = run_local_agent_dry_run(
        root,
        host=StrictSyntheticAgentHost(),
        use_existing_synthetic_source=True,
    )

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.SOURCE_UNSAFE"
    assert _tree_hashes(root / "runtime") == runtime_before


def test_acceptance_gate_rejects_missing_disabled_ocr_failure_and_cli_is_nonzero(
    tmp_path, monkeypatch, capsys
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    monkeypatch.setattr(
        DataCleaningTools,
        "prepare_file_organization_run",
        lambda _self, _paths: {"status": "success", "failures": []},
    )
    from scripts.run_local_agent_dry_run import main

    root = tmp_path / "dry-run"
    code = main(["--root", str(root)])
    payload = json.loads(capsys.readouterr().out)

    assert code != 0
    assert payload["status"] == "blocked"
    assert payload["error_code"] == "DRY_RUN.ACCEPTANCE_FAILED"


class SourceMutatingHost(StrictSyntheticAgentHost):
    def __init__(self, source_path: Path) -> None:
        super().__init__()
        self.source_path = source_path

    def write_responses(self, request_path: Path, response_path: Path) -> None:
        super().write_responses(request_path, response_path)
        self.source_path.write_text("mutated after hand-off", encoding="utf-8")


def test_acceptance_gate_rejects_source_hash_change(tmp_path, monkeypatch) -> None:
    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / "dry-run"
    result = run_local_agent_dry_run(
        root,
        host=SourceMutatingHost(root / "source" / "synthetic-governance.md"),
    )

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.ACCEPTANCE_FAILED"
    assert result["hashes_before"] != result["hashes_after"]


def test_acceptance_gate_rejects_any_archive_result(tmp_path, monkeypatch) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.execute_archive_plan
    root = tmp_path / "dry-run"

    def plant_archive_result(self, run_id, confirmed=False):
        result = original(self, run_id, confirmed=confirmed)
        artifact = root / "runtime" / "runs" / run_id / "archive_result.json"
        artifact.write_text('{"executed": false}', encoding="utf-8")
        return result

    monkeypatch.setattr(
        DataCleaningTools, "execute_archive_plan", plant_archive_result
    )
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    result = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.ACCEPTANCE_FAILED"


@pytest.mark.parametrize(
    "method_name",
    [
        "verify_file_organization_run",
        "audit_file_organization_run",
        "prepare_feedback_form",
    ],
)
def test_acceptance_gate_requires_successful_post_run_checks(
    tmp_path, monkeypatch, method_name
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    monkeypatch.setattr(
        DataCleaningTools,
        method_name,
        lambda _self, _run_id: {"status": "failed"},
    )
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    result = run_local_agent_dry_run(
        tmp_path / method_name, host=StrictSyntheticAgentHost()
    )

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.ACCEPTANCE_FAILED"


@pytest.mark.parametrize(
    "counterexample",
    ["native_status", "invoice_candidate", "review_item", "confirmed_action"],
)
def test_acceptance_gate_rejects_invalid_review_only_result(
    tmp_path, monkeypatch, counterexample
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.resume_agent_judgement_run

    def mutate_result(self, run_id, response_path):
        result = original(self, run_id, response_path)
        if counterexample == "native_status":
            result["candidate_interpretations"][0]["status"] = "success"
        elif counterexample == "invoice_candidate":
            invoice = next(
                item
                for item in result["candidate_interpretations"]
                if item["source_ref"]["object_key"] == "synthetic-invoice.pdf"
            )
            invoice["business_relation"]["candidate_contract_id"] = "C-999"
        elif counterexample == "review_item":
            result["review_queue"]["items"][0]["question"] = ""
        else:
            result["archive_actions"][0]["confirmed"] = True
        return result

    monkeypatch.setattr(
        DataCleaningTools, "resume_agent_judgement_run", mutate_result
    )
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    result = run_local_agent_dry_run(
        tmp_path / counterexample, host=StrictSyntheticAgentHost()
    )

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.ACCEPTANCE_FAILED"


def test_scan_fixture_is_a_raster_image_only_pdf(tmp_path, monkeypatch) -> None:
    import fitz

    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / "dry-run"
    result = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())

    assert result["status"] == "needs_review"
    with fitz.open(root / "source" / "synthetic-scan.pdf") as scan:
        assert scan.page_count == 1
        page = scan[0]
        assert page.get_text("text").strip() == ""
        assert page.get_images(full=True)
        assert page.get_drawings() == []


def test_unexpected_cli_failure_redacts_root_from_stderr(
    tmp_path, monkeypatch, capsys
) -> None:
    import scripts.run_local_agent_dry_run as module

    root = tmp_path / "private-dry-run-root"

    def fail_with_sensitive_path(*_args, **_kwargs):
        raise OSError(f"cannot write {root}")

    monkeypatch.setattr(module, "run_local_agent_dry_run", fail_with_sensitive_path)

    code = module.main(["--root", str(root)])
    captured = capsys.readouterr()

    assert code != 0
    assert str(root) not in captured.err
    assert json.loads(captured.err)["error_code"] == "DRY_RUN.UNEXPECTED"


def test_independent_archive_binding_e2e_remains_review_only(
    tmp_path, monkeypatch
) -> None:
    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / "dry-run"
    result = run_local_agent_dry_run(
        root,
        host=StrictSyntheticAgentHost(),
        include_archive_binding=True,
    )

    assert result["status"] == "needs_review"
    config = json.loads(
        (root / "config" / "project-manager.local.json").read_text(
            encoding="utf-8"
        )
    )
    bindings = {item["binding_id"]: item for item in config["storage_bindings"]}
    assert set(bindings) == {"dry-run-source", "dry-run-archive"}
    archive_binding = bindings["dry-run-archive"]
    archive_root = Path(archive_binding["physical_root"]).resolve()
    assert archive_binding["roles"] == ["archive_target"]
    assert archive_binding["readable"] is False
    assert archive_binding["writable"] is True
    assert archive_root == (root / "archive").resolve()
    assert archive_root != (root / "source").resolve()
    assert archive_root != (root / "runtime").resolve()
    assert list(archive_root.iterdir()) == []
    plan = json.loads(
        next((root / "runtime" / "runs").glob("*/planned_archive_actions.json"))
        .read_text(encoding="utf-8")
    )
    assert plan["actions"]
    assert all(action["confirmed"] is False for action in plan["actions"])
    assert not list((root / "runtime").rglob("archive_result.json"))


def _rewrite_self_authored_marker_hash(source_file: Path) -> None:
    import hashlib

    marker = source_file.parent / ".local-agent-dry-run.synthetic.v1.json"
    payload = json.loads(marker.read_text(encoding="utf-8"))
    payload["hashes"][source_file.name] = hashlib.sha256(
        source_file.read_bytes()
    ).hexdigest()
    marker.write_text(json.dumps(payload), encoding="utf-8")


def test_reuse_rejects_invoice_pdf_vector_drawing_with_forged_marker(
    tmp_path, monkeypatch
) -> None:
    """A PDF drawing outside normalized text must invalidate source reuse."""
    import fitz

    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / "dry-run"
    first = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())
    assert first["status"] == "needs_review"
    runtime_before = _tree_hashes(root / "runtime")

    invoice = root / "source" / "synthetic-invoice.pdf"
    forged = invoice.with_name("synthetic-invoice.forged.pdf")
    with fitz.open(invoice) as document:
        document[0].draw_rect((24, 24, 48, 48), color=(1, 0, 0))
        document.save(forged)
    os.replace(forged, invoice)
    _rewrite_self_authored_marker_hash(invoice)

    result = run_local_agent_dry_run(
        root,
        host=StrictSyntheticAgentHost(),
        use_existing_synthetic_source=True,
    )

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.SOURCE_UNSAFE"
    assert _tree_hashes(root / "runtime") == runtime_before


@pytest.mark.parametrize("format_name", ["docx_header", "xlsx_hyperlink"])
def test_reuse_rejects_off_contract_package_structure_with_forged_marker(
    tmp_path, monkeypatch, format_name
) -> None:
    """Package relationships and non-cell content are part of canonical reuse."""
    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / format_name
    first = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())
    assert first["status"] == "needs_review"
    runtime_before = _tree_hashes(root / "runtime")

    if format_name == "docx_header":
        from docx import Document

        source_file = root / "source" / "synthetic-contract.docx"
        document = Document(source_file)
        document.sections[0].header.paragraphs[0].text = (
            "synthetic off-contract header"
        )
        document.save(source_file)
    else:
        from openpyxl import load_workbook

        source_file = root / "source" / "synthetic-project.xlsx"
        workbook = load_workbook(source_file)
        workbook["synthetic-project"]["A1"].hyperlink = (
            "https://example.invalid/synthetic-link"
        )
        workbook.save(source_file)
        workbook.close()
    _rewrite_self_authored_marker_hash(source_file)

    result = run_local_agent_dry_run(
        root,
        host=StrictSyntheticAgentHost(),
        use_existing_synthetic_source=True,
    )

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.SOURCE_UNSAFE"
    assert _tree_hashes(root / "runtime") == runtime_before


def _assert_cli_acceptance_failure(root: Path, capsys) -> None:
    from scripts.run_local_agent_dry_run import main

    code = main(["--root", str(root)])
    captured = capsys.readouterr()
    assert code == 2, captured.err
    payload = json.loads(captured.out)
    assert payload["status"] == "blocked"
    assert payload["error_code"] == "DRY_RUN.ACCEPTANCE_FAILED"


@pytest.mark.parametrize(
    "counterexample",
    ["corrupt_verdict", "needs_correction", "wrong_run"],
)
def test_acceptance_gate_rejects_invalid_or_unbound_verification_artifact(
    tmp_path, monkeypatch, capsys, counterexample
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.verify_file_organization_run

    def mutate_verification(self, run_id):
        result = original(self, run_id)
        if counterexample == "wrong_run":
            result["run_id"] = "run_wrong_verification"
        else:
            result["overall_verdict"] = (
                "corrupt"
                if counterexample == "corrupt_verdict"
                else "needs_correction"
            )
        Path(result["artifact_path"]).write_text(
            json.dumps(result), encoding="utf-8"
        )
        return result

    monkeypatch.setattr(
        DataCleaningTools, "verify_file_organization_run", mutate_verification
    )

    _assert_cli_acceptance_failure(tmp_path / counterexample, capsys)


@pytest.mark.parametrize("counterexample", ["minimal", "wrong_run"])
def test_acceptance_gate_rejects_invalid_or_unbound_audit_artifact(
    tmp_path, monkeypatch, capsys, counterexample
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.audit_file_organization_run

    def mutate_audit(self, run_id):
        result = original(self, run_id)
        if counterexample == "minimal":
            return {"status": "success"}
        result["run_id"] = "run_wrong_audit"
        Path(result["artifact_path"]).write_text(
            json.dumps(result), encoding="utf-8"
        )
        return result

    monkeypatch.setattr(
        DataCleaningTools, "audit_file_organization_run", mutate_audit
    )

    _assert_cli_acceptance_failure(tmp_path / counterexample, capsys)


@pytest.mark.parametrize("counterexample", ["minimal", "wrong_run"])
def test_acceptance_gate_rejects_invalid_or_unbound_feedback_artifact(
    tmp_path, monkeypatch, capsys, counterexample
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.prepare_feedback_form

    def mutate_feedback(self, run_id):
        result = original(self, run_id)
        if counterexample == "minimal":
            return {"status": "success"}
        result["run_id"] = "run_wrong_feedback"
        form_path = Path(result["artifacts"]["feedback_form_json"])
        form = json.loads(form_path.read_text(encoding="utf-8"))
        form["run_id"] = "run_wrong_feedback"
        form_path.write_text(json.dumps(form), encoding="utf-8")
        return result

    monkeypatch.setattr(
        DataCleaningTools, "prepare_feedback_form", mutate_feedback
    )

    _assert_cli_acceptance_failure(tmp_path / counterexample, capsys)


@pytest.mark.parametrize(
    "counterexample", ["fake_item", "duplicate_item", "wrong_run"]
)
def test_acceptance_gate_rejects_invalid_or_unbound_review_artifact(
    tmp_path, monkeypatch, capsys, counterexample
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.resume_agent_judgement_run

    def mutate_review(self, run_id, response_path):
        result = original(self, run_id, response_path)
        queue = result["review_queue"]
        if counterexample == "fake_item":
            queue["items"][0] = {
                "id": "R-FAKE",
                "question": "synthetic fake review item",
            }
        elif counterexample == "duplicate_item":
            queue["items"][1]["id"] = queue["items"][0]["id"]
        else:
            queue["run_id"] = "run_wrong_review"
        Path(result["artifacts"]["review_queue"]).write_text(
            json.dumps(queue), encoding="utf-8"
        )
        return result

    monkeypatch.setattr(
        DataCleaningTools, "resume_agent_judgement_run", mutate_review
    )

    _assert_cli_acceptance_failure(tmp_path / counterexample, capsys)


@pytest.mark.parametrize(
    ("source_name", "filename"),
    [
        ("pdf", "synthetic-invoice.pdf"),
        ("docx", "synthetic-contract.docx"),
        ("xlsx", "synthetic-project.xlsx"),
    ],
)
def test_reuse_rejects_trailing_container_payload_with_forged_marker(
    tmp_path, monkeypatch, source_name, filename
) -> None:
    """Raw bytes after a valid container terminator are still source content."""
    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / source_name
    first = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())
    assert first["status"] == "needs_review"
    runtime_before = _tree_hashes(root / "runtime")

    source_file = root / "source" / filename
    with source_file.open("ab") as stream:
        stream.write(b"\nsynthetic-trailing-container-payload\n")
    _rewrite_self_authored_marker_hash(source_file)

    result = run_local_agent_dry_run(
        root,
        host=StrictSyntheticAgentHost(),
        use_existing_synthetic_source=True,
    )

    assert result["status"] == "blocked"
    assert result["error_code"] == "DRY_RUN.SOURCE_UNSAFE"
    assert _tree_hashes(root / "runtime") == runtime_before


def test_legal_reuse_does_not_allocate_a_system_temporary_path(
    tmp_path, monkeypatch
) -> None:
    """Canonical reuse must not write a reference fixture outside supplied root."""
    import tempfile

    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / "dry-run"
    first = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())
    assert first["status"] == "needs_review"

    names = (
        "TemporaryDirectory",
        "NamedTemporaryFile",
        "TemporaryFile",
        "mkdtemp",
        "mkstemp",
    )
    originals = {name: getattr(tempfile, name) for name in names}

    def root_local_temp(name):
        def guarded(*args, **kwargs):
            raw_dir = kwargs.get("dir")
            if raw_dir is None or not Path(raw_dir).resolve().is_relative_to(
                root.resolve()
            ):
                raise AssertionError(
                    "system temporary paths are outside the supplied root"
                )
            return originals[name](*args, **kwargs)

        return guarded

    for name in names:
        monkeypatch.setattr(tempfile, name, root_local_temp(name))

    result = run_local_agent_dry_run(
        root,
        host=StrictSyntheticAgentHost(),
        use_existing_synthetic_source=True,
    )

    assert result["status"] == "needs_review"


def test_acceptance_gate_rejects_schema_valid_forged_review_question(
    tmp_path, monkeypatch, capsys
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.resume_agent_judgement_run

    def mutate_question(self, run_id, response_path):
        result = original(self, run_id, response_path)
        result["review_queue"]["items"][0]["question"] = (
            "Synthetic schema-valid question supplied by an untrusted stage?"
        )
        Path(result["artifacts"]["review_queue"]).write_text(
            json.dumps(result["review_queue"]), encoding="utf-8"
        )
        return result

    monkeypatch.setattr(
        DataCleaningTools, "resume_agent_judgement_run", mutate_question
    )

    _assert_cli_acceptance_failure(tmp_path / "forged-question", capsys)


def test_acceptance_gate_rejects_high_risk_finding_with_pass_verdict(
    tmp_path, monkeypatch, capsys
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.verify_file_organization_run

    def mutate_verification(self, run_id):
        result = original(self, run_id)
        result["findings"].append(
            {
                "id": "AV-FORGED-HIGH",
                "dimension": "field_completeness",
                "severity": "high",
                "confidence": 0.99,
                "file": "synthetic-invoice.pdf",
                "message": "Synthetic high-risk finding hidden by a pass verdict.",
                "details": {"missing_fields": ["synthetic_required_field"]},
            }
        )
        result["finding_count"] = len(result["findings"])
        result["overall_verdict"] = "pass"
        result["needs_human_review"] = False
        result["archive_allowed"] = True
        result["next_actions"] = ["audit_file_organization_run"]
        Path(result["artifact_path"]).write_text(
            json.dumps(result), encoding="utf-8"
        )
        return result

    monkeypatch.setattr(
        DataCleaningTools, "verify_file_organization_run", mutate_verification
    )

    _assert_cli_acceptance_failure(tmp_path / "high-finding-pass", capsys)


def test_acceptance_gate_rejects_audit_violation_and_forged_trace(
    tmp_path, monkeypatch, capsys
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.audit_file_organization_run

    def mutate_audit(self, run_id):
        result = original(self, run_id)
        result["policy_violations"] = [
            {
                "id": "AUD-FORGED-HIGH",
                "severity": "high",
                "rule": "synthetic_forged_rule",
                "message": "Synthetic high policy violation.",
            }
        ]
        result["loop_trace_summary"] = {
            "trace_id": "forged-trace",
            "status": "success",
            "round_count": 9,
        }
        persisted = dict(result)
        artifact_path = Path(persisted.pop("artifact_path"))
        artifact_path.write_text(json.dumps(persisted), encoding="utf-8")
        return result

    monkeypatch.setattr(DataCleaningTools, "audit_file_organization_run", mutate_audit)

    _assert_cli_acceptance_failure(tmp_path / "forged-audit", capsys)


@pytest.mark.parametrize(
    ("field", "forged_value"),
    [
        ("risk_level", "P0"),
        ("source_file", "synthetic-forged-source.pdf"),
        ("target_path", "synthetic-forged-target"),
        ("field", "synthetic_forged_field"),
        ("expected_field", "synthetic_forged_expected_field"),
    ],
)
def test_acceptance_gate_rejects_forged_feedback_projection(
    tmp_path, monkeypatch, capsys, field, forged_value
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.prepare_feedback_form

    def mutate_feedback(self, run_id):
        result = original(self, run_id)
        form_path = Path(result["artifacts"]["feedback_form_json"])
        form = json.loads(form_path.read_text(encoding="utf-8"))
        form["items"][0][field] = forged_value
        form_path.write_text(json.dumps(form), encoding="utf-8")
        return result

    monkeypatch.setattr(DataCleaningTools, "prepare_feedback_form", mutate_feedback)

    _assert_cli_acceptance_failure(tmp_path / field, capsys)


def test_acceptance_gate_rejects_forged_feedback_markdown(
    tmp_path, monkeypatch, capsys
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.prepare_feedback_form

    def mutate_markdown(self, run_id):
        result = original(self, run_id)
        markdown_path = Path(result["artifacts"]["feedback_form_md"])
        with markdown_path.open("a", encoding="utf-8") as stream:
            stream.write("\nSynthetic forged Markdown payload.\n")
        return result

    monkeypatch.setattr(DataCleaningTools, "prepare_feedback_form", mutate_markdown)

    _assert_cli_acceptance_failure(tmp_path / "forged-markdown", capsys)


def test_acceptance_gate_rejects_duplicate_archive_action_coverage(
    tmp_path, monkeypatch, capsys
) -> None:
    import copy

    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.resume_agent_judgement_run

    def duplicate_archive_action(self, run_id, response_path):
        result = original(self, run_id, response_path)
        result["archive_actions"][1] = copy.deepcopy(result["archive_actions"][0])
        plan_path = Path(result["artifacts"]["planned_archive_actions"])
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        plan["actions"] = result["archive_actions"]
        plan_path.write_text(json.dumps(plan), encoding="utf-8")
        return result

    monkeypatch.setattr(
        DataCleaningTools,
        "resume_agent_judgement_run",
        duplicate_archive_action,
    )

    _assert_cli_acceptance_failure(tmp_path / "duplicate-archive", capsys)


def test_acceptance_gate_rejects_fake_archive_execution_return(
    tmp_path, monkeypatch, capsys
) -> None:
    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)

    def fake_archive_return(_self, _run_id, confirmed=False):
        assert confirmed is False
        return {
            "schema_version": "synthetic-fake-archive.v1",
            "status": "success",
            "run_id": "run_wrong_archive",
            "confirmed": False,
        }

    monkeypatch.setattr(
        DataCleaningTools, "execute_archive_plan", fake_archive_return
    )

    _assert_cli_acceptance_failure(tmp_path / "fake-archive-return", capsys)


def test_canonical_source_digest_generation_is_fully_in_memory(monkeypatch) -> None:
    import tempfile

    import scripts.run_local_agent_dry_run as module

    def forbidden(*_args, **_kwargs):
        raise AssertionError("canonical reference generation wrote to a filesystem path")

    for name in ("mkdir", "write_text", "write_bytes", "unlink", "rmdir"):
        monkeypatch.setattr(module.Path, name, forbidden)
    monkeypatch.setattr(module.os, "replace", forbidden)
    for name in (
        "TemporaryDirectory",
        "NamedTemporaryFile",
        "TemporaryFile",
        "mkdtemp",
        "mkstemp",
    ):
        monkeypatch.setattr(tempfile, name, forbidden)

    digests = module._trusted_canonical_source_digests()

    assert set(digests) == set(module.SOURCE_NAMES)
    assert all(
        len(digest) == 64 and set(digest) <= set("0123456789abcdef")
        for digest in digests.values()
    )


def test_two_legal_reuses_and_forged_retries_do_not_pollute_root(
    tmp_path, monkeypatch
) -> None:
    _forbid_external_providers(monkeypatch)
    from scripts.run_local_agent_dry_run import run_local_agent_dry_run

    root = tmp_path / "dry-run"
    first = run_local_agent_dry_run(root, host=StrictSyntheticAgentHost())
    second = run_local_agent_dry_run(
        root, host=StrictSyntheticAgentHost(), use_existing_synthetic_source=True
    )
    third = run_local_agent_dry_run(
        root, host=StrictSyntheticAgentHost(), use_existing_synthetic_source=True
    )

    assert first["status"] == second["status"] == third["status"] == "needs_review"
    source_file = root / "source" / "synthetic-invoice.pdf"
    with source_file.open("ab") as stream:
        stream.write(b"\nsynthetic-forged-retry\n")
    _rewrite_self_authored_marker_hash(source_file)
    runtime_before = _tree_hashes(root / "runtime")
    children_before = sorted(path.name for path in root.iterdir())

    for _ in range(2):
        blocked = run_local_agent_dry_run(
            root,
            host=StrictSyntheticAgentHost(),
            use_existing_synthetic_source=True,
        )
        assert blocked["status"] == "blocked"
        assert blocked["error_code"] == "DRY_RUN.SOURCE_UNSAFE"

    assert _tree_hashes(root / "runtime") == runtime_before
    assert sorted(path.name for path in root.iterdir()) == children_before
    assert not (root / ".local-agent-dry-run-canonical-reference").exists()


def test_strict_markdown_read_rejects_path_replacement_after_lstat(
    tmp_path, monkeypatch
) -> None:
    import scripts.run_local_agent_dry_run as module

    target = tmp_path / "feedback.md"
    replacement = tmp_path / "replacement.md"
    target.write_text("expected", encoding="utf-8")
    replacement.write_text("forged", encoding="utf-8")
    original_lstat = module.os.lstat
    replaced = False

    def replace_after_lstat(path):
        nonlocal replaced
        details = original_lstat(path)
        if Path(path) == target and not replaced:
            os.replace(replacement, target)
            replaced = True
        return details

    monkeypatch.setattr(module.os, "lstat", replace_after_lstat)

    with pytest.raises(ValueError, match="unsafe text artifact"):
        module._strict_text_artifact(target)


def test_acceptance_rejects_forged_pass_over_strict_missing_candidate_fields(
    tmp_path, monkeypatch, capsys
) -> None:
    import copy

    from tools.data_cleaning_tools import DataCleaningTools

    _forbid_external_providers(monkeypatch)
    original = DataCleaningTools.verify_file_organization_run
    observed = {}

    def remove_required_fields_then_forge_pass(self, run_id):
        extracted_dir = Path(self.workspace_dir) / "runs" / run_id / "extracted"
        contract_path = next(
            path
            for path in extracted_dir.glob("*.json")
            if json.loads(path.read_text(encoding="utf-8"))["document_type"]
            == "合同"
        )
        artifact = json.loads(contract_path.read_text(encoding="utf-8"))
        artifact["candidate_fields"] = {}
        contract_path.write_text(
            json.dumps(artifact, ensure_ascii=False), encoding="utf-8"
        )
        result = original(self, run_id)
        observed.update(copy.deepcopy(result))
        result["findings"] = []
        result["finding_count"] = 0
        result["overall_verdict"] = "pass"
        result["needs_human_review"] = False
        result["archive_allowed"] = True
        result["next_actions"] = ["audit_file_organization_run"]
        Path(result["artifact_path"]).write_text(
            json.dumps(result), encoding="utf-8"
        )
        return result

    monkeypatch.setattr(
        DataCleaningTools,
        "verify_file_organization_run",
        remove_required_fields_then_forge_pass,
    )

    _assert_cli_acceptance_failure(tmp_path / "strict-missing-fields", capsys)
    assert observed["overall_verdict"] == "needs_correction"
    assert any(finding["severity"] == "high" for finding in observed["findings"])
