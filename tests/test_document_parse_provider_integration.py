"""Provider-selection and runtime-cache tests for document_parse KB lookup."""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from docx import Document

from app_bootstrap import composition as bootstrap_composition
from integrations.pageindex import structure_index as pageindex_adapter_module
from integrations.pageindex.pageindex_client import PageIndexClient
from platform_core.models import (
    CapabilityReport,
    StructureIndexRequest,
    StructureIndexResult,
)
from platform_core.settings import load_app_settings
from skills.document_parse import kb
from skills.document_parse.parse import parse


def _write_kb_document(path: Path) -> None:
    path.write_text(
        "# 文档分类\n\n## 招标公告\n\n招标公告与采购公告。\n",
        encoding="utf-8",
    )


def _write_pageindex_config(
    path: Path,
    *,
    pageindex_dir: Path,
    runtime_workspace: Path,
) -> None:
    path.write_text(
        json.dumps(
            {
                "deployment_mode": "local",
                "business_root": str(path.parent / "business"),
                "runtime_workspace": str(runtime_workspace),
                "providers": {
                    "document_store": "local",
                    "structure_index": "pageindex",
                    "pageindex_dir": str(pageindex_dir),
                    "projection_writer": "filesystem",
                    "projection_root": str(runtime_workspace / "projections"),
                },
            }
        ),
        encoding="utf-8",
    )


def _install_fake_pageindex(pageindex_dir: Path) -> None:
    client = PageIndexClient(str(pageindex_dir))
    Path(client.python_bin).parent.mkdir(parents=True)
    Path(client.python_bin).touch()
    Path(client.cli_script).touch()


def _pageindex_runner(actual_cwds):
    def run(command, **kwargs):
        if command[-1] == "--version":
            return subprocess.CompletedProcess(command, 0, "Python 3.12", "")
        cwd = Path(kwargs["cwd"])
        actual_cwds.append(cwd)
        staged_source = Path(command[command.index("--md_path") + 1])
        output = cwd / "results" / f"{staged_source.stem}_structure.json"
        output.parent.mkdir(parents=True)
        output.write_text(
            json.dumps(
                {
                    "doc_name": staged_source.name,
                    "structure": [
                        {
                            "title": "文档分类",
                            "node_id": "0000",
                            "start_index": 1,
                            "end_index": 3,
                            "nodes": [
                                {
                                    "title": "招标公告",
                                    "node_id": "0001",
                                    "start_index": 2,
                                    "end_index": 2,
                                    "nodes": [],
                                },
                                {
                                    "title": "采购公告",
                                    "node_id": "0002",
                                    "start_index": 3,
                                    "end_index": 3,
                                    "nodes": [],
                                },
                            ],
                        }
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        return subprocess.CompletedProcess(command, 0, "", "")

    return run


def test_disabled_provider_never_constructs_pageindex_client(
    tmp_path, monkeypatch
):
    kb_document = tmp_path / "kb.md"
    _write_kb_document(kb_document)
    monkeypatch.setattr(kb, "KB_DOC_PATH", str(kb_document))
    monkeypatch.setenv("PROJECT_MANAGER_STRUCTURE_INDEX", "disabled")
    monkeypatch.setenv(
        "PROJECT_MANAGER_PAGEINDEX_DIR",
        str(tmp_path / "Users" / "alice" / "PageIndex-secret"),
    )
    monkeypatch.setenv("PROJECT_MANAGER_WORKSPACE_DIR", str(tmp_path / "runtime"))
    constructor_calls = []

    class ForbiddenPageIndexClient:
        def __init__(self, *args, **kwargs):
            constructor_calls.append((args, kwargs))
            raise AssertionError("disabled provider must not construct PageIndex")

    class ForbiddenPageIndexStructureIndex:
        def __init__(self, *args, **kwargs):
            constructor_calls.append((args, kwargs))
            raise AssertionError("disabled provider must not construct adapter")

    monkeypatch.setattr(
        pageindex_adapter_module,
        "PageIndexClient",
        ForbiddenPageIndexClient,
    )
    monkeypatch.setattr(
        bootstrap_composition,
        "PageIndexStructureIndex",
        ForbiddenPageIndexStructureIndex,
    )

    result = kb.query_kb(
        {"raw_text": "普通文档"},
        source_path="ordinary.docx",
        force_reindex=True,
    )

    assert result is None
    assert constructor_calls == []


def test_json_config_selects_pageindex_and_writes_cache_under_runtime_workspace(
    tmp_path, monkeypatch
):
    kb_document = tmp_path / "kb.md"
    _write_kb_document(kb_document)
    monkeypatch.setattr(kb, "KB_DOC_PATH", str(kb_document))
    pageindex_dir = tmp_path / "configured-pageindex"
    runtime_workspace = tmp_path / "runtime"
    config_file = tmp_path / "project-manager.local.json"
    _install_fake_pageindex(pageindex_dir)
    _write_pageindex_config(
        config_file,
        pageindex_dir=pageindex_dir,
        runtime_workspace=runtime_workspace,
    )
    actual_cwds = []
    monkeypatch.setattr(subprocess, "run", _pageindex_runner(actual_cwds))

    result = kb.get_kb_structure(
        force_reindex=True,
        config_file=config_file,
        environ={},
    )

    cache_path = runtime_workspace / "cache" / "document_parse" / "kb_cache.json"
    assert result["status"] == "success"
    assert result["engine"] == "pageindex"
    assert cache_path.is_file()
    assert not (Path(kb.__file__).resolve().parent / "kb_cache.json").exists()
    assert len(actual_cwds) == 1
    assert actual_cwds[0].is_relative_to(runtime_workspace / "pageindex" / "staging")
    assert actual_cwds[0] != pageindex_dir


def test_environment_disabled_overrides_json_pageindex_without_subprocess(
    tmp_path, monkeypatch
):
    kb_document = tmp_path / "kb.md"
    _write_kb_document(kb_document)
    monkeypatch.setattr(kb, "KB_DOC_PATH", str(kb_document))
    pageindex_dir = tmp_path / "Users" / "alice" / "PageIndex-secret"
    runtime_workspace = tmp_path / "runtime-secret"
    config_file = tmp_path / "project-manager.local.json"
    _install_fake_pageindex(pageindex_dir)
    _write_pageindex_config(
        config_file,
        pageindex_dir=pageindex_dir,
        runtime_workspace=runtime_workspace,
    )
    subprocess_calls = []
    monkeypatch.setattr(
        subprocess,
        "run",
        lambda *args, **kwargs: subprocess_calls.append((args, kwargs)),
    )
    diagnostics = []

    result = kb.query_kb(
        {"raw_text": "普通文档"},
        source_path="ordinary.docx",
        force_reindex=True,
        config_file=config_file,
        environ={"PROJECT_MANAGER_STRUCTURE_INDEX": "disabled"},
        diagnostics=diagnostics,
    )

    assert result is None
    assert subprocess_calls == []
    assert diagnostics == [
        {
            "component": "document_parse.knowledge_base",
            "status": "degraded",
            "error_code": "DOCUMENT_PARSE.KB.STRUCTURE_INDEX_DISABLED",
            "provider": "disabled",
            "message": "Knowledge-base structure indexing is disabled.",
        }
    ]
    serialized = json.dumps(diagnostics, ensure_ascii=False)
    for secret in (str(pageindex_dir), str(runtime_workspace), "alice", "secret"):
        assert secret not in serialized


def test_injected_structure_index_receives_content_identity_and_cache_path(
    tmp_path, monkeypatch
):
    kb_document = tmp_path / "kb.md"
    _write_kb_document(kb_document)
    monkeypatch.setattr(kb, "KB_DOC_PATH", str(kb_document))
    cache_path = tmp_path / "injected-cache" / "kb.json"

    class RecordingStructureIndex:
        name = "recording"
        provider_version = "recording-v3"

        def __init__(self):
            self.requests = []

        def probe(self):
            return CapabilityReport("ready", self.name, self.provider_version, "ready")

        def index(self, request: StructureIndexRequest):
            self.requests.append(request)
            return StructureIndexResult(
                "success",
                self.name,
                "recording://artifact-one",
                (
                    {
                        "title": "文档分类",
                        "node_id": "0000",
                        "nodes": [
                            {
                                "title": "招标公告",
                                "node_id": "0001",
                                "nodes": [],
                            }
                        ],
                    },
                ),
                "",
                "",
            )

    adapter = RecordingStructureIndex()

    invalid_config = tmp_path / "invalid-config.json"
    invalid_config.write_text("{not-json", encoding="utf-8")

    result = kb.get_kb_structure(
        force_reindex=True,
        structure_index=adapter,
        cache_path=cache_path,
        managed_cache_root=cache_path.parent,
        config_file=invalid_config,
        environ={
            "PROJECT_MANAGER_DEPLOYMENT_MODE": "central",
            "PROJECT_MANAGER_DATABASE_PROVIDER": "sqlite",
        },
    )

    expected_hash = hashlib.sha256(kb_document.read_bytes()).hexdigest()
    assert result["status"] == "success"
    assert cache_path.is_file()
    assert len(adapter.requests) == 1
    request = adapter.requests[0]
    assert request.document_version_id.startswith("document-parse-kb:")
    assert request.content_hash == expected_hash
    assert request.source_path == str(kb_document)
    assert request.media_type == "text/markdown"


def test_kb_cache_identity_uses_content_hash_and_provider_version(
    tmp_path, monkeypatch
):
    kb_document = tmp_path / "kb.md"
    _write_kb_document(kb_document)
    monkeypatch.setattr(kb, "KB_DOC_PATH", str(kb_document))
    cache_path = tmp_path / "cache" / "kb.json"

    class VersionedStructureIndex:
        name = "recording"

        def __init__(self):
            self.provider_version = "recording-v1"
            self.requests = []

        def probe(self):
            return CapabilityReport(
                "ready",
                self.name,
                self.provider_version,
                "ready",
            )

        def index(self, request: StructureIndexRequest):
            self.requests.append(request)
            return StructureIndexResult(
                "success",
                self.name,
                f"recording://artifact-{len(self.requests)}",
                (
                    {
                        "title": "文档分类",
                        "node_id": "0000",
                        "nodes": [
                            {
                                "title": "招标公告",
                                "node_id": "0001",
                                "nodes": [],
                            }
                        ],
                    },
                ),
                "",
                "",
            )

    adapter = VersionedStructureIndex()
    first = kb.get_kb_structure(
        force_reindex=True,
        structure_index=adapter,
        cache_path=cache_path,
        managed_cache_root=cache_path.parent,
    )
    cached = kb.get_kb_structure(
        structure_index=adapter,
        cache_path=cache_path,
        managed_cache_root=cache_path.parent,
    )

    assert first == cached
    assert len(adapter.requests) == 1

    original_stat = kb_document.stat()
    original_bytes = kb_document.read_bytes()
    changed_bytes = original_bytes.replace(
        "采购".encode("utf-8"),
        "征集".encode("utf-8"),
    )
    assert len(changed_bytes) == len(original_bytes)
    kb_document.write_bytes(changed_bytes)
    os.utime(
        kb_document,
        ns=(original_stat.st_atime_ns, original_stat.st_mtime_ns),
    )

    kb.get_kb_structure(
        structure_index=adapter,
        cache_path=cache_path,
        managed_cache_root=cache_path.parent,
    )

    assert len(adapter.requests) == 2
    assert adapter.requests[0].content_hash != adapter.requests[1].content_hash

    adapter.provider_version = "recording-v2"
    kb.get_kb_structure(
        structure_index=adapter,
        cache_path=cache_path,
        managed_cache_root=cache_path.parent,
    )

    assert len(adapter.requests) == 3
    assert adapter.requests[1].content_hash == adapter.requests[2].content_hash


def test_concurrent_cache_writes_preserve_each_provider_identity(
    tmp_path, monkeypatch
):
    kb_document = tmp_path / "kb.md"
    _write_kb_document(kb_document)
    monkeypatch.setattr(kb, "KB_DOC_PATH", str(kb_document))
    cache_path = tmp_path / "cache" / "kb.json"
    rendezvous = threading.Barrier(2)

    class ConcurrentStructureIndex:
        name = "recording"

        def __init__(self, provider_version):
            self.provider_version = provider_version
            self.requests = []

        def probe(self):
            return CapabilityReport(
                "ready",
                self.name,
                self.provider_version,
                "ready",
            )

        def index(self, request):
            self.requests.append(request)
            if len(self.requests) == 1:
                rendezvous.wait(timeout=5)
            return StructureIndexResult(
                "success",
                self.name,
                f"recording://{self.provider_version}",
                (
                    {
                        "title": "文档分类",
                        "node_id": "0000",
                        "nodes": [
                            {
                                "title": self.provider_version,
                                "node_id": "0001",
                                "nodes": [],
                            }
                        ],
                    },
                ),
                "",
                "",
            )

    first = ConcurrentStructureIndex("recording-v1")
    second = ConcurrentStructureIndex("recording-v2")

    with ThreadPoolExecutor(max_workers=2) as pool:
        first_future = pool.submit(
            kb.get_kb_structure,
            True,
            structure_index=first,
            cache_path=cache_path,
            managed_cache_root=cache_path.parent,
        )
        second_future = pool.submit(
            kb.get_kb_structure,
            True,
            structure_index=second,
            cache_path=cache_path,
            managed_cache_root=cache_path.parent,
        )
        first_future.result(timeout=10)
        second_future.result(timeout=10)

    assert len(first.requests) == len(second.requests) == 1

    kb.get_kb_structure(structure_index=first, cache_path=cache_path, managed_cache_root=cache_path.parent)
    kb.get_kb_structure(structure_index=second, cache_path=cache_path, managed_cache_root=cache_path.parent)

    assert len(first.requests) == len(second.requests) == 1


def test_missing_provider_version_falls_back_without_index_or_cache(
    tmp_path, monkeypatch
):
    kb_document = tmp_path / "kb.md"
    _write_kb_document(kb_document)
    monkeypatch.setattr(kb, "KB_DOC_PATH", str(kb_document))
    cache_path = tmp_path / "cache" / "kb.json"
    diagnostics = []

    class MissingIdentityStructureIndex:
        name = "anonymous"
        provider_version = None

        def __init__(self):
            self.index_calls = 0

        def probe(self):
            raise RuntimeError("provider exposes no stable identity")

        def index(self, request):
            self.index_calls += 1
            raise AssertionError("provider without identity must not index")

    adapter = MissingIdentityStructureIndex()

    result = kb.query_kb(
        {"raw_text": "普通文档"},
        structure_index=adapter,
        cache_path=cache_path,
        managed_cache_root=cache_path.parent,
        diagnostics=diagnostics,
    )

    assert result is None
    assert adapter.index_calls == 0
    assert not cache_path.exists()
    assert diagnostics[0]["error_code"] == (
        "DOCUMENT_PARSE.KB.PROVIDER_IDENTITY_UNAVAILABLE"
    )


def test_unreadable_pageindex_manifest_never_indexes_or_writes_kb_cache(
    tmp_path, monkeypatch
):
    kb_document = tmp_path / "kb.md"
    _write_kb_document(kb_document)
    monkeypatch.setattr(kb, "KB_DOC_PATH", str(kb_document))
    pageindex_dir = tmp_path / "pageindex"
    cache_path = tmp_path / "cache" / "kb.json"
    diagnostics = []
    subprocess_calls = []

    def unavailable_manifest(self):
        raise OSError("provider manifest cannot be read")

    monkeypatch.setattr(
        PageIndexClient,
        "_provider_manifest_digest",
        unavailable_manifest,
    )
    monkeypatch.setattr(
        subprocess,
        "run",
        lambda *args, **kwargs: subprocess_calls.append((args, kwargs)),
    )
    adapter = pageindex_adapter_module.PageIndexStructureIndex(pageindex_dir)

    result = kb.query_kb(
        {"raw_text": "普通文档"},
        structure_index=adapter,
        cache_path=cache_path,
        managed_cache_root=cache_path.parent,
        diagnostics=diagnostics,
    )

    assert result is None
    assert subprocess_calls == []
    assert not cache_path.exists()
    assert diagnostics[0]["error_code"] == (
        "DOCUMENT_PARSE.KB.PROVIDER_IDENTITY_UNAVAILABLE"
    )


def test_provider_version_change_during_index_is_not_cached(
    tmp_path, monkeypatch
):
    kb_document = tmp_path / "kb.md"
    _write_kb_document(kb_document)
    monkeypatch.setattr(kb, "KB_DOC_PATH", str(kb_document))
    cache_path = tmp_path / "cache" / "kb.json"
    diagnostics = []

    class MutatingStructureIndex:
        name = "recording"

        def __init__(self):
            self.provider_version = "recording-v1"
            self.index_calls = 0

        def index(self, request):
            self.index_calls += 1
            self.provider_version = "recording-v2"
            return StructureIndexResult(
                "success",
                self.name,
                "recording://unstable",
                (
                    {
                        "title": "文档分类",
                        "node_id": "0000",
                        "nodes": [
                            {
                                "title": "招标公告",
                                "node_id": "0001",
                                "nodes": [],
                            }
                        ],
                    },
                ),
                "",
                "",
            )

    adapter = MutatingStructureIndex()
    result = kb.query_kb(
        {"raw_text": "普通文档"},
        structure_index=adapter,
        cache_path=cache_path,
        managed_cache_root=cache_path.parent,
        diagnostics=diagnostics,
    )

    assert result is None
    assert adapter.index_calls == 1
    assert not cache_path.exists()
    assert diagnostics[0]["error_code"] == (
        "DOCUMENT_PARSE.KB.PROVIDER_IDENTITY_CHANGED"
    )


def test_parse_v1_default_shape_keeps_diagnostics_opt_in(
    tmp_path, monkeypatch
):
    kb_document = tmp_path / "kb.md"
    _write_kb_document(kb_document)
    monkeypatch.setattr(kb, "KB_DOC_PATH", str(kb_document))
    document_path = tmp_path / "notice.docx"
    document = Document()
    document.add_paragraph("招标公告测试")
    document.save(document_path)
    settings = load_app_settings(
        config_file="",
        environ={
            "PROJECT_MANAGER_BUSINESS_ROOT": str(tmp_path / "business"),
            "PROJECT_MANAGER_WORKSPACE_DIR": str(tmp_path / "runtime"),
            "PROJECT_MANAGER_STRUCTURE_INDEX": "disabled",
        },
    )

    result = parse(str(document_path), app_settings=settings)

    assert result["schema_version"] == "document_parse.v1"
    assert "diagnostics" not in result


def test_parse_disabled_provider_falls_back_with_structured_sanitized_diagnostic(
    tmp_path, monkeypatch
):
    kb_document = tmp_path / "kb.md"
    _write_kb_document(kb_document)
    monkeypatch.setattr(kb, "KB_DOC_PATH", str(kb_document))
    document_path = tmp_path / "notice.docx"
    document = Document()
    document.add_paragraph("招标公告测试")
    document.save(document_path)
    secret_pageindex_dir = tmp_path / "Users" / "alice" / "PageIndex-secret"
    settings = load_app_settings(
        config_file="",
        environ={
            "PROJECT_MANAGER_BUSINESS_ROOT": str(tmp_path / "business"),
            "PROJECT_MANAGER_WORKSPACE_DIR": str(tmp_path / "runtime-secret"),
            "PROJECT_MANAGER_STRUCTURE_INDEX": "disabled",
            "PROJECT_MANAGER_PAGEINDEX_DIR": str(secret_pageindex_dir),
        },
    )

    result = parse(
        str(document_path),
        app_settings=settings,
        include_diagnostics=True,
    )

    assert result["status"] == "success"
    assert result["business_judgement"]["category"] == "招标公告"
    assert result["knowledge_base_used"] is False
    assert result["diagnostics"][0]["error_code"] == (
        "DOCUMENT_PARSE.KB.STRUCTURE_INDEX_DISABLED"
    )
    serialized = json.dumps(result["diagnostics"], ensure_ascii=False)
    for secret in (str(secret_pageindex_dir), "alice", "PageIndex-secret"):
        assert secret not in serialized
