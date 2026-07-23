"""document_parse skill 端到端测试。"""

import os
import tempfile
import unittest

from docx import Document
import openpyxl

from skills.document_parse import parse, validate, SCHEMA_VERSION
from skills.document_parse.router import route, NoExecutorError
from skills.document_parse.validator import validate_or_raise, DocumentParseValidationError


class TestDocumentParseSkill(unittest.TestCase):

    def setUp(self):
        self.docx_path = None
        self.xlsx_path = None

    def tearDown(self):
        for p in [self.docx_path, self.xlsx_path]:
            if p and os.path.exists(p):
                os.unlink(p)

    def _create_docx(self, paragraphs):
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            self.docx_path = f.name
        doc = Document()
        for p in paragraphs:
            doc.add_paragraph(p)
        doc.save(self.docx_path)

    def _create_xlsx(self, rows):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            self.xlsx_path = f.name
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in rows:
            ws.append(row)
        wb.save(self.xlsx_path)

    # ---- 端到端: docx 分类 ----

    def test_parse_docx_zhaobiao_gonggao(self):
        self._create_docx(["招标公告测试", "项目名称:测试A", "采购人:客户B"])
        result = parse(self.docx_path)
        self.assertEqual(result["schema_version"], SCHEMA_VERSION)
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["executor_used"], "docx")
        self.assertEqual(result["business_judgement"]["category"], "招标公告")
        self.assertEqual(result["business_judgement"]["confidence"], "high")
        self.assertEqual(result["validation"]["schema_valid"], True)
        self.assertGreater(len(result["extracted_data"]["paragraphs"]), 0)

    def test_parse_docx_caigou_gonggao(self):
        self._create_docx(["采购公告", "项目名称:测试B"])
        result = parse(self.docx_path)
        self.assertEqual(result["business_judgement"]["category"], "招标公告")

    def test_parse_docx_toubiao(self):
        self._create_docx(["投标文件响应", "投标人:公司X", "报价:95万"])
        result = parse(self.docx_path)
        self.assertEqual(result["business_judgement"]["category"], "投标文件")

    def test_parse_docx_hetong(self):
        self._create_docx(["合同文本", "签订合同", "甲方:测试A", "乙方:测试B"])
        result = parse(self.docx_path)
        self.assertEqual(result["business_judgement"]["category"], "合同文件")

    def test_parse_docx_baoming(self):
        self._create_docx(["报名表", "项目名称:测试A"])
        result = parse(self.docx_path)
        self.assertEqual(result["business_judgement"]["category"], "报名材料")

    def test_parse_docx_other(self):
        self._create_docx(["普通文档", "没有关键词"])
        result = parse(self.docx_path)
        self.assertEqual(result["business_judgement"]["category"], "其他")

    # ---- 端到端: xlsx ----

    def test_parse_xlsx_success(self):
        self._create_xlsx([["项目", "金额"], ["测试A", "100万"]])
        result = parse(self.xlsx_path)
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["executor_used"], "xlsx")
        self.assertGreater(len(result["extracted_data"]["sheets"]), 0)

    # ---- 端到端: URL stub ----

    def test_parse_url_stub_blocked(self):
        result = parse("https://example.com/page.html")
        self.assertEqual(result["status"], "blocked")
        self.assertEqual(result["executor_used"], "url")
        self.assertEqual(result["executor_implementation_status"], "stub")
        self.assertEqual(result["reason"], "executor_not_implemented")

    def test_parse_url_ftp_scheme_blocked(self):
        """ftp:// 不被 router 识别为 URL(只接 http/https),走 NoExecutorError。"""
        result = parse("ftp://example.com")
        self.assertEqual(result["status"], "blocked")
        self.assertEqual(result["reason"], "parse_error")
        self.assertIn("No executor can handle", result["validation"]["warnings"][0])

    def test_parse_url_https_extracts_to_stub(self):
        """https:// 路由到 UrlExecutor stub(返回 blocked + executor_not_implemented)。"""
        result = parse("https://example.com")
        self.assertEqual(result["status"], "blocked")
        self.assertEqual(result["executor_used"], "url")
        self.assertEqual(result["reason"], "executor_not_implemented")

    # ---- 端到端: 错误处理 ----

    def test_parse_unsupported_type(self):
        result = parse("/tmp/test.unknown")
        self.assertEqual(result["status"], "blocked")
        self.assertEqual(result["reason"], "parse_error")

    def test_parse_docx_file_not_found(self):
        result = parse("/nonexistent.docx")
        self.assertEqual(result["status"], "blocked")

    # ---- schema 校验 ----

    def test_success_results_pass_schema(self):
        self._create_docx(["招标公告"])
        result = parse(self.docx_path)
        ok, errs = validate(result)
        self.assertTrue(ok, f"Schema invalid: {errs}")

    def test_blocked_results_pass_schema(self):
        """blocked 结果也必须 schema 有效(_build_blocked_result 保证)。"""
        result = parse("/nonexistent.docx")
        ok, errs = validate(result)
        self.assertTrue(ok, f"Blocked result schema invalid: {errs}")

    def test_url_stub_result_passes_schema(self):
        result = parse("https://example.com")
        ok, errs = validate(result)
        self.assertTrue(ok, f"Stub result schema invalid: {errs}")

    def test_validate_or_raise_raises_on_invalid(self):
        with self.assertRaises(DocumentParseValidationError):
            validate_or_raise({"schema_version": "wrong"})

    def test_validate_or_raise_passes_on_valid(self):
        result = parse("/nonexistent.docx")  # blocked 但 schema 有效
        validate_or_raise(result)  # 不抛

    # ---- run_id ----

    def test_run_id_auto_generated(self):
        self._create_docx(["test"])
        result = parse(self.docx_path)
        self.assertTrue(result["run_id"].startswith("run-"))

    def test_run_id_custom_provided(self):
        self._create_docx(["test"])
        custom_id = "run-my-custom-123"
        result = parse(self.docx_path, run_id=custom_id)
        self.assertEqual(result["run_id"], custom_id)

    # ---- elapsed_seconds ----

    def test_elapsed_seconds_recorded(self):
        self._create_docx(["test"])
        result = parse(self.docx_path)
        self.assertGreater(result["elapsed_seconds"], 0)


class TestRouter(unittest.TestCase):

    def test_route_docx(self):
        executor, file_type = route("/tmp/a.docx")
        self.assertEqual(executor.name, "docx")
        self.assertEqual(file_type, ".docx")

    def test_route_xlsx(self):
        executor, file_type = route("/tmp/a.xlsx")
        self.assertEqual(executor.name, "xlsx")

    def test_route_url(self):
        executor, file_type = route("https://example.com/page")
        self.assertEqual(executor.name, "url")
        self.assertEqual(file_type, "url")

    def test_route_unsupported_raises(self):
        with self.assertRaises(NoExecutorError):
            route("/tmp/a.unknown")
