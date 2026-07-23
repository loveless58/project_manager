"""Docx (.docx) 解析执行器。纯技术,无业务规则。

输出 raw_data 包含:paragraphs (非空段落), tables (rows of cells), metadata (author/title/created/modified), raw_text (段落 + 表格串联)。
"""

import os
from typing import Any, Dict, List

from integrations.executors.base import ParseExecutor


class DocxExecutor(ParseExecutor):
    """Word (.docx) 解析器。"""

    name = "docx"

    def can_handle(self, source: str) -> bool:
        if self._is_url(source):
            return False
        return source.lower().endswith(".docx")

    def extract(self, source: str) -> Dict[str, Any]:
        try:
            self._check_source_exists(source)
        except FileNotFoundError as e:
            return self._blocked(str(e))

        try:
            from docx import Document  # type: ignore
        except ImportError:
            return self._blocked(
                "python-docx not installed. Run: pip install python-docx"
            )

        try:
            doc = Document(source)

            # 段落(非空)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

            # 表格(过滤全空行)
            tables = []
            for table in doc.tables:
                rows = []
                for row in table.rows:
                    cells = [cell.text for cell in row.cells]
                    if any(c.strip() for c in cells):
                        rows.append(cells)
                if rows:
                    tables.append(rows)

            # metadata
            metadata: Dict[str, str] = {}
            cp = doc.core_properties
            if cp:
                metadata = {
                    "author": cp.author or "",
                    "title": cp.title or "",
                    "subject": cp.subject or "",
                    "created": str(cp.created) if cp.created else "",
                    "modified": str(cp.modified) if cp.modified else "",
                }

            # raw_text(段落 + 表格串联)
            parts = list(paragraphs)
            for table in tables:
                for row in table:
                    parts.append(" | ".join(row))
            raw_text = "\n".join(parts)

            return {
                "status": "success",
                "raw_data": {
                    "paragraphs": paragraphs,
                    "tables": tables,
                    "metadata": metadata,
                    "raw_text": raw_text,
                },
                "error": None,
                "implementation_status": "implemented",
            }
        except Exception as e:
            return self._blocked(f"docx parse error: {type(e).__name__}: {e}")

    def get_supported_types(self) -> List[str]:
        return [".docx"]
