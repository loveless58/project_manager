#!/usr/bin/env python3
"""Run a synthetic, local-only agent judgement acceptance cycle.

The runner creates every business-like input itself under an explicit root.  It
does not discover user files, construct a configured model adapter, probe OCR
engines, or authorize an archive action.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import stat
import sys
from typing import Any, Protocol, Sequence


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from app_bootstrap.composition import build_runtime_adapters
from infrastructure.database.sqlite import initialize_schema_metadata
from integrations.business_context import JsonBusinessContextProvider
from ocr.providers import DisabledOcrProvider
from platform_core import load_app_settings
from platform_core.path_locality import is_obvious_network_location
from services.agent_judgement_gateway import AgentJudgementGateway
from services.retrieval_service import RetrievalService
from tools.data_cleaning_tools import DataCleaningTools


SCHEMA_VERSION = "local_agent_judgement_dry_run.v1"
SOURCE_BINDING_ID = "dry-run-source"
ARCHIVE_BINDING_ID = "dry-run-archive"
CONFIG_TEMPLATE = PROJECT_ROOT / "config" / "local-agent-dry-run.example.json"
CATALOG_TEMPLATE = (
    PROJECT_ROOT
    / "tests"
    / "fixtures"
    / "agent_judgement_business_context_catalog.v1.json"
)
SOURCE_NAMES = {
    "invoice": "synthetic-invoice.pdf",
    "contract": "synthetic-contract.docx",
    "markdown": "synthetic-governance.md",
    "xlsx": "synthetic-project.xlsx",
    "xml": "synthetic-project.xml",
    "scan": "synthetic-scan.pdf",
}
_SYNTHETIC_MARKER = ".local-agent-dry-run.synthetic.v1.json"
_INVOICE_CODE = "".join(("INV", "-001"))
_CONTRACT_CODE = "".join(("CT", "-001"))
_PROJECT_CODE = "".join(("PRJ", "-001"))
_INVOICE_TEXT = (
    f"电子发票\n发票号码：{_INVOICE_CODE}\n开票日期：2026-07-26\n"
    "购买方名称：合成甲方有限公司\n购买方税号：913100000000000001\n"
    "销售方名称：合成乙方有限公司\n销售方税号：913100000000000002\n"
    "项目名称 | 规格型号 | 金额\n技术服务 | 合成版 | 100.00\n" + "A" * 160
)
_CONTRACT_PARAGRAPHS = (
    "合成技术服务合同",
    f"合同登记编号：{_CONTRACT_CODE}",
    f"项目编号：{_PROJECT_CODE}",
)
_MARKDOWN_TEXT = (
    f"# Synthetic governance\n\n项目编号：{_PROJECT_CODE}\n"
    f"合同编号：{_CONTRACT_CODE}\n"
)
_XLSX_ROWS = ((f"项目编号：{_PROJECT_CODE}",), (f"合同编号：{_CONTRACT_CODE}",))
_XML_TEXT = (
    "<?xml version=\"1.0\" encoding=\"UTF-8\"?>"
    f"<project><项目编号>{_PROJECT_CODE}</项目编号>"
    f"<合同编号>{_CONTRACT_CODE}</合同编号></project>"
)


class AgentHost(Protocol):
    """Narrow host boundary: read one request artifact and write one response."""

    def write_responses(self, request_path: Path, response_path: Path) -> None:
        ...


class DryRunSafetyError(ValueError):
    """An explicit dry-run filesystem boundary is unsafe."""

    def __init__(self, code: str) -> None:
        self.code = code
        super().__init__(code)


class DeterministicSyntheticAgentHost:
    """Offline host used only for the script's generated synthetic documents."""

    def write_responses(self, request_path: Path, response_path: Path) -> None:
        request_artifact = json.loads(request_path.read_text(encoding="utf-8"))
        responses = [
            self._response(request) for request in request_artifact["requests"]
        ]
        response_path.parent.mkdir(parents=True, exist_ok=True)
        response_path.write_text(
            json.dumps(
                {
                    "schema_version": "agent_judgement_responses.v1",
                    "run_id": request_artifact["run_id"],
                    "responses": responses,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

    @staticmethod
    def _response(request: dict[str, Any]) -> dict[str, Any]:
        interpretation_request = request["interpretation_request"]
        document_type = interpretation_request["document"]["document_type_hint"]
        evidence = interpretation_request["business_context"]["evidence"][:1]
        return {
            "schema_version": "agent_judgement_response.v1",
            "run_id": request["run_id"],
            "request_id": request["request_id"],
            "request_hash": request["request_hash"],
            "interpreter": "strict_synthetic_agent",
            "model": "deterministic_fixture_v1",
            "created_at": "2026-07-26T00:00:00Z",
            "interpretation": {
                "schema_version": "candidate_document_interpretation.v1",
                "status": "success",
                "document_type": document_type,
                "fields": {},
                "relations": [
                    {
                        "relation_type": (
                            "invoice_contract"
                            if document_type == "invoice"
                            else "contract_project"
                        ),
                        "target_candidate_id": "C-001",
                    }
                ],
                "evidence": evidence,
                "confidence": 0.94,
                "interpreter": "strict_synthetic_agent",
                "model": "deterministic_fixture_v1",
                "prompt_version": "document_interpretation.v1",
                "policy_version": "document_interpretation_policy.v1",
            },
        }


def _is_relative_to(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
    except ValueError:
        return False
    return True


def _is_reparse_point(path: Path) -> bool:
    try:
        details = os.lstat(path)
    except FileNotFoundError:
        return False
    except OSError:
        raise DryRunSafetyError("DRY_RUN.PATH_UNSAFE") from None
    attributes = getattr(details, "st_file_attributes", 0)
    reparse_flag = getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0x400)
    return stat.S_ISLNK(details.st_mode) or bool(attributes & reparse_flag)


def _reject_reparse_components(path: Path) -> None:
    current = Path(path.anchor)
    for part in path.parts[1:]:
        current /= part
        if _is_reparse_point(current):
            raise DryRunSafetyError("DRY_RUN.PATH_REPARSE")


def _resolved_fixed_children(raw_root: Path, resolved_root: Path) -> dict[str, Path]:
    children: dict[str, Path] = {}
    for name in ("source", "runtime", "config", "archive"):
        raw_child = raw_root / name
        _reject_reparse_components(raw_child)
        resolved_child = raw_child.resolve(strict=False)
        if resolved_child == resolved_root or not _is_relative_to(
            resolved_child, resolved_root
        ):
            raise DryRunSafetyError("DRY_RUN.PATH_OUT_OF_SCOPE")
        children[name] = resolved_child
    return children


def _safe_root(root: os.PathLike[str] | str) -> Path:
    try:
        raw_value = os.fspath(root)
    except TypeError:
        raise DryRunSafetyError("DRY_RUN.ROOT_UNSAFE") from None
    if type(raw_value) is not str or not raw_value.strip():
        raise DryRunSafetyError("DRY_RUN.ROOT_UNSAFE")
    if is_obvious_network_location(raw_value):
        raise DryRunSafetyError("DRY_RUN.ROOT_NETWORK")
    raw = Path(os.path.normpath(raw_value))
    if not raw.is_absolute():
        raise DryRunSafetyError("DRY_RUN.ROOT_MUST_BE_ABSOLUTE")
    _reject_reparse_components(raw)
    resolved = raw.resolve(strict=False)
    _resolved_fixed_children(raw, resolved)
    dangerous = {
        Path(resolved.anchor).resolve(),
        Path.home().resolve(),
        PROJECT_ROOT.resolve(),
    }
    if resolved in dangerous or (resolved / ".git").exists():
        raise DryRunSafetyError("DRY_RUN.ROOT_UNSAFE")
    if resolved.exists() and not resolved.is_dir():
        raise DryRunSafetyError("DRY_RUN.ROOT_UNSAFE")
    return resolved


def _validated_source_paths(source_root: Path, paths: Sequence[Path]) -> list[Path]:
    validated: list[Path] = []
    for raw_path in paths:
        path = Path(raw_path).resolve()
        if (
            not _is_relative_to(path, source_root)
            or not path.is_file()
            or path.is_symlink()
        ):
            raise DryRunSafetyError("SOURCE_BINDING.OUT_OF_SCOPE")
        validated.append(path)
    return validated


def _hashes(paths: Sequence[Path]) -> dict[str, str]:
    return {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in paths
    }


def _write_pdf(path: Path, text: str = "", *, image_only: bool = False) -> None:
    import fitz

    document = fitz.open()
    page = document.new_page()
    if image_only:
        pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 64, 64), False)
        pixmap.clear_with(115)
        page.insert_image(page.rect, pixmap=pixmap)
    else:
        import html

        page.insert_htmlbox(
            page.rect,
            f"<p>{html.escape(text).replace(chr(10), '<br>')}</p>",
        )
    document.save(path)
    document.close()


def _create_synthetic_source(source_root: Path) -> dict[str, Path]:
    from docx import Document
    from openpyxl import Workbook

    paths = {name: source_root / filename for name, filename in SOURCE_NAMES.items()}
    if any(path.exists() for path in paths.values()):
        raise DryRunSafetyError("DRY_RUN.SOURCE_ALREADY_EXISTS")

    _write_pdf(
        paths["invoice"],
        _INVOICE_TEXT,
    )

    contract = Document()
    for paragraph in _CONTRACT_PARAGRAPHS:
        contract.add_paragraph(paragraph)
    contract.save(paths["contract"])

    paths["markdown"].write_text(_MARKDOWN_TEXT, encoding="utf-8")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "synthetic-project"
    for row in _XLSX_ROWS:
        worksheet.append(row)
    workbook.save(paths["xlsx"])

    paths["xml"].write_text(_XML_TEXT, encoding="utf-8")
    _write_pdf(paths["scan"], image_only=True)
    return paths


def _has_canonical_synthetic_content(paths: dict[str, Path]) -> bool:
    """Validate normalized fixture semantics against constants owned by this code."""
    try:
        import fitz
        from docx import Document
        from openpyxl import load_workbook
        from xml.etree import ElementTree

        if paths["markdown"].read_text(encoding="utf-8") != _MARKDOWN_TEXT:
            return False
        if paths["xml"].read_text(encoding="utf-8") != _XML_TEXT:
            return False
        xml_root = ElementTree.parse(paths["xml"]).getroot()
        if (
            xml_root.tag != "project"
            or xml_root.attrib
            or [(child.tag, child.text, child.attrib) for child in xml_root]
            != [
                ("项目编号", _PROJECT_CODE, {}),
                ("合同编号", _CONTRACT_CODE, {}),
            ]
        ):
            return False

        contract = Document(paths["contract"])
        if tuple(item.text for item in contract.paragraphs) != _CONTRACT_PARAGRAPHS:
            return False
        if contract.tables or contract.inline_shapes:
            return False

        workbook = load_workbook(paths["xlsx"], read_only=True, data_only=False)
        try:
            if workbook.sheetnames != ["synthetic-project"]:
                return False
            rows = tuple(
                tuple(cell.value for cell in row)
                for row in workbook["synthetic-project"].iter_rows()
            )
            if rows != _XLSX_ROWS or workbook._external_links:
                return False
        finally:
            workbook.close()

        with fitz.open(paths["invoice"]) as invoice:
            if invoice.page_count != 1 or invoice.embfile_count() != 0:
                return False
            page = invoice[0]
            normalized_text = "".join(page.get_text("text").split())
            if normalized_text != "".join(_INVOICE_TEXT.split()):
                return False
            if page.get_images(full=True):
                return False

        with fitz.open(paths["scan"]) as scan:
            if scan.page_count != 1 or scan.embfile_count() != 0:
                return False
            page = scan[0]
            images = page.get_images(full=True)
            if page.get_text("text").strip() or page.get_drawings():
                return False
            if len(images) != 1:
                return False
            xref, _smask, width, height, bpc = images[0][:5]
            if (width, height, bpc) != (64, 64, 8):
                return False
            image_rects = page.get_image_rects(xref)
            if len(image_rects) != 1 or image_rects[0] != page.rect:
                return False
            pixels = fitz.Pixmap(scan, xref)
            try:
                if (
                    pixels.n != 3
                    or pixels.alpha != 0
                    or pixels.samples != bytes([115]) * (64 * 64 * 3)
                ):
                    return False
            finally:
                pixels = None
    except (OSError, ValueError, KeyError, TypeError):
        return False
    return True


def _load_existing_synthetic_source(source_root: Path) -> dict[str, Path]:
    expected = {name: source_root / filename for name, filename in SOURCE_NAMES.items()}
    marker = source_root / _SYNTHETIC_MARKER
    allowed = {path.name for path in expected.values()} | {_SYNTHETIC_MARKER}
    if any(path.name not in allowed for path in source_root.iterdir()):
        raise DryRunSafetyError("DRY_RUN.SOURCE_UNSAFE")
    if not marker.is_file() or marker.is_symlink():
        raise DryRunSafetyError("DRY_RUN.SOURCE_UNSAFE")
    try:
        recorded = json.loads(marker.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        raise DryRunSafetyError("DRY_RUN.SOURCE_UNSAFE") from None
    files = _validated_source_paths(source_root, list(expected.values()))
    if not _has_canonical_synthetic_content(expected):
        raise DryRunSafetyError("DRY_RUN.SOURCE_UNSAFE")
    if recorded != {
        "schema_version": "local_agent_dry_run_source.v1",
        "hashes": _hashes(files),
    }:
        raise DryRunSafetyError("DRY_RUN.SOURCE_UNSAFE")
    return expected


def _prepare_source(
    source_root: Path,
    *,
    use_existing_synthetic_source: bool,
) -> dict[str, Path]:
    if source_root.exists():
        if source_root.is_symlink() or not source_root.is_dir():
            raise DryRunSafetyError("DRY_RUN.SOURCE_UNSAFE")
        if not use_existing_synthetic_source:
            raise DryRunSafetyError("DRY_RUN.SOURCE_ALREADY_EXISTS")
        if not any(source_root.iterdir()):
            paths = _create_synthetic_source(source_root)
        else:
            return _load_existing_synthetic_source(source_root)
    else:
        source_root.mkdir(parents=True)
        paths = _create_synthetic_source(source_root)
    marker = source_root / _SYNTHETIC_MARKER
    files = list(paths.values())
    marker.write_text(
        json.dumps(
            {
                "schema_version": "local_agent_dry_run_source.v1",
                "hashes": _hashes(files),
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return paths


def _replace_tokens(value: Any, replacements: dict[str, str]) -> Any:
    if isinstance(value, dict):
        return {key: _replace_tokens(item, replacements) for key, item in value.items()}
    if isinstance(value, list):
        return [_replace_tokens(item, replacements) for item in value]
    if isinstance(value, str) and value in replacements:
        return replacements[value]
    return value


def _write_runtime_config(
    *,
    root: Path,
    source_root: Path,
    runtime_root: Path,
    config_root: Path,
    include_archive_binding: bool,
) -> Path:
    template = json.loads(CONFIG_TEMPLATE.read_text(encoding="utf-8"))
    sqlite_path = runtime_root / "state" / "project_manager.sqlite3"
    projection_root = runtime_root / "projections"
    payload = _replace_tokens(
        template,
        {
            "__DRY_RUN_SOURCE_ROOT__": str(source_root),
            "__DRY_RUN_RUNTIME_ROOT__": str(runtime_root),
            "__DRY_RUN_SQLITE_PATH__": str(sqlite_path),
            "__DRY_RUN_PROJECTION_ROOT__": str(projection_root),
        },
    )
    if include_archive_binding:
        archive_root = root / "archive"
        archive_root.mkdir(exist_ok=True)
        payload["storage_bindings"].append(
            {
                "binding_id": ARCHIVE_BINDING_ID,
                "provider": "local",
                "node_id": "dry-run-node",
                "logical_root": "business://dry-run-archive/",
                "physical_root": str(archive_root),
                "roles": ["archive_target"],
                "readable": False,
                "writable": True,
                "enabled": True,
            }
        )
    config_path = config_root / "project-manager.local.json"
    config_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return config_path


def _write_runtime_catalog(config_root: Path, context_document: Path) -> Path:
    payload = json.loads(CATALOG_TEMPLATE.read_text(encoding="utf-8"))
    document = payload["records"][0]["documents"][0]
    document["path"] = str(context_document)
    document["content_hash"] = hashlib.sha256(context_document.read_bytes()).hexdigest()
    catalog_path = config_root / "business-context-catalog.json"
    catalog_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return catalog_path


def _blocked(
    code: str,
    *,
    hashes_before: dict[str, str] | None = None,
    hashes_after: dict[str, str] | None = None,
) -> dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION,
        "status": "blocked",
        "error_code": code,
        "hashes_before": hashes_before or {},
        "hashes_after": hashes_after or {},
        "archive_execution": {"status": "not_run", "confirmed": False},
    }


def _confirmation_flags_are_false(*payloads: Any) -> bool:
    flags: list[Any] = []

    def collect(value: Any) -> None:
        if isinstance(value, dict):
            for key, item in value.items():
                if key == "confirmed":
                    flags.append(item)
                collect(item)
        elif isinstance(value, list):
            for item in value:
                collect(item)

    for payload in payloads:
        collect(payload)
    return all(flag is False for flag in flags)


def _acceptance_succeeded(
    *,
    native_statuses: dict[str, str],
    scan_failure: str,
    invoice_candidate: str,
    hashes_before: dict[str, str],
    hashes_after: dict[str, str],
    review_items: list[dict[str, Any]],
    verification: dict[str, Any],
    audit: dict[str, Any],
    feedback: dict[str, Any],
    resumed: dict[str, Any],
    archive: dict[str, Any],
    runtime_root: Path,
    run_id: str,
) -> bool:
    expected_statuses = {
        label: "needs_review"
        for label in ("invoice", "contract", "markdown", "xlsx", "xml")
    }
    if native_statuses != expected_statuses:
        return False
    if scan_failure != "OCR.CAPABILITY_DISABLED":
        return False
    if invoice_candidate != "C-001" or hashes_before != hashes_after:
        return False
    if not review_items or not all(
        isinstance(item, dict)
        and isinstance(item.get("id"), str)
        and bool(item["id"].strip())
        and isinstance(item.get("question"), str)
        and bool(item["question"].strip())
        for item in review_items
    ):
        return False
    if (
        verification.get("schema_version") != "adversarial_verification.v1"
        or not verification.get("overall_verdict")
        or audit.get("status") != "success"
        or feedback.get("status") != "success"
    ):
        return False

    returned_actions = resumed.get("archive_actions")
    if not isinstance(returned_actions, list) or not returned_actions:
        return False
    plan_path = runtime_root / "runs" / run_id / "planned_archive_actions.json"
    try:
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return False
    planned_actions = plan.get("actions") if isinstance(plan, dict) else None
    if not isinstance(planned_actions, list) or not planned_actions:
        return False
    if not all(action.get("confirmed") is False for action in returned_actions):
        return False
    if not all(action.get("confirmed") is False for action in planned_actions):
        return False
    if not _confirmation_flags_are_false(resumed, plan, archive):
        return False
    return not any(runtime_root.rglob("archive_result.json"))


def run_local_agent_dry_run(
    root: os.PathLike[str] | str,
    *,
    host: AgentHost | None,
    use_existing_synthetic_source: bool = False,
    include_archive_binding: bool = False,
    source_paths: Sequence[os.PathLike[str] | str] | None = None,
) -> dict[str, Any]:
    """Run one deterministic, synthetic review-only cycle under ``root``."""
    try:
        root_path = _safe_root(root)
        if root_path.exists():
            existing_names = {item.name for item in root_path.iterdir()}
            allowed_names = {"source", "runtime", "config", "archive"}
            if existing_names and (
                "source" not in existing_names
                or existing_names - allowed_names
            ):
                raise DryRunSafetyError("DRY_RUN.ROOT_NOT_EMPTY")
        source_root = root_path / "source"
        if source_paths is not None:
            _validated_source_paths(
                source_root.resolve(), [Path(path) for path in source_paths]
            )
        root_path.mkdir(parents=True, exist_ok=True)
        paths = _prepare_source(
            source_root,
            use_existing_synthetic_source=use_existing_synthetic_source,
        )
        all_sources = _validated_source_paths(source_root, list(paths.values()))
    except DryRunSafetyError as exc:
        return _blocked(exc.code)

    before = _hashes(all_sources)
    runtime_root = root_path / "runtime"
    config_root = root_path / "config"
    runtime_root.mkdir(exist_ok=True)
    config_root.mkdir(exist_ok=True)
    config_path = _write_runtime_config(
        root=root_path,
        source_root=source_root,
        runtime_root=runtime_root,
        config_root=config_root,
        include_archive_binding=include_archive_binding,
    )
    catalog_path = _write_runtime_catalog(config_root, paths["markdown"])

    settings = load_app_settings(config_file=config_path, environ={})
    sqlite_path = settings.database.sqlite_path
    if sqlite_path is None or not _is_relative_to(sqlite_path, runtime_root):
        return _blocked(
            "DRY_RUN.SQLITE_OUT_OF_SCOPE",
            hashes_before=before,
            hashes_after=_hashes(all_sources),
        )
    sqlite_path.parent.mkdir(parents=True, exist_ok=True)
    initialize_schema_metadata(sqlite_path)

    adapters = build_runtime_adapters(settings)
    registry = adapters.storage_binding_registry
    if registry is None:
        return _blocked("SOURCE_BINDING.INVALID", hashes_before=before)
    context = JsonBusinessContextProvider(catalog_path, storage_bindings=registry)
    retrieval = RetrievalService(context, adapters.structure_index)
    tools = DataCleaningTools(
        workspace_dir=str(runtime_root),
        ocr_adapter=DisabledOcrProvider().extract,
        storage_binding_registry=registry,
        document_store_router=adapters.document_store_router,
        retrieval_service=retrieval,
        archive_target_resolver=adapters.archive_target_resolver,
        agent_judgement_gateway=AgentJudgementGateway(),
    )

    scan_result = tools.prepare_file_organization_run([str(paths["scan"])])
    scan_failures = scan_result.get("failures") or []
    scan_failure = (
        str(scan_failures[0].get("blocked_reason", "DOCUMENT_PARSE.FAILED"))
        if scan_failures
        else ""
    )

    native_labels = ("invoice", "contract", "markdown", "xlsx", "xml")
    native_paths = [str(paths[label]) for label in native_labels]
    prepared = tools.prepare_agent_judgement_run(native_paths)
    if prepared.get("status") != "awaiting_agent_judgement":
        return _blocked(
            str(prepared.get("blocked_reason", "AGENT_JUDGEMENT.REQUEST_INVALID")),
            hashes_before=before,
            hashes_after=_hashes(all_sources),
        )

    request_path = Path(prepared["artifacts"]["agent_judgement_requests"])
    request_digest = hashlib.sha256(request_path.read_bytes()).hexdigest()
    response_path = runtime_root / "agent-host-responses" / f"{prepared['run_id']}.json"
    if host is not None:
        response_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            host.write_responses(request_path, response_path)
        except Exception:
            return _blocked(
                "AGENT_JUDGEMENT.HOST_FAILED",
                hashes_before=before,
                hashes_after=_hashes(all_sources),
            )
    if hashlib.sha256(request_path.read_bytes()).hexdigest() != request_digest:
        return _blocked(
            "AGENT_JUDGEMENT.REQUEST_TAMPERED",
            hashes_before=before,
            hashes_after=_hashes(all_sources),
        )

    resumed = tools.resume_agent_judgement_run(
        prepared["run_id"], str(response_path)
    )
    if resumed.get("status") != "success":
        return _blocked(
            str(resumed.get("blocked_reason", "AGENT_JUDGEMENT.RESPONSE_INVALID")),
            hashes_before=before,
            hashes_after=_hashes(all_sources),
        )

    run_id = resumed["run_id"]
    verification = tools.verify_file_organization_run(run_id)
    audit = tools.audit_file_organization_run(run_id)
    feedback = tools.prepare_feedback_form(run_id)
    archive = tools.execute_archive_plan(run_id, confirmed=False)
    after = _hashes(all_sources)

    label_by_name = {paths[label].name: label for label in native_labels}
    native_statuses: dict[str, str] = {}
    invoice_candidate = ""
    for item in resumed["candidate_interpretations"]:
        label = label_by_name[item["source_ref"]["object_key"]]
        native_statuses[label] = item["status"]
        if label == "invoice":
            invoice_candidate = item["business_relation"].get(
                "candidate_contract_id", ""
            )

    review_items = resumed["review_queue"]["items"]
    if not _acceptance_succeeded(
        native_statuses=native_statuses,
        scan_failure=scan_failure,
        invoice_candidate=invoice_candidate,
        hashes_before=before,
        hashes_after=after,
        review_items=review_items,
        verification=verification,
        audit=audit,
        feedback=feedback,
        resumed=resumed,
        archive=archive,
        runtime_root=runtime_root,
        run_id=run_id,
    ):
        return _blocked(
            "DRY_RUN.ACCEPTANCE_FAILED",
            hashes_before=before,
            hashes_after=after,
        )
    return {
        "schema_version": SCHEMA_VERSION,
        "status": "needs_review",
        "run_id": run_id,
        "native_statuses": native_statuses,
        "scan_failure": scan_failure,
        "invoice_candidate": invoice_candidate,
        "hashes_before": before,
        "hashes_after": after,
        "review_items_valid": [
            bool(item.get("id") and item.get("question")) for item in review_items
        ],
        "verification_status": verification.get("status")
        or verification.get("overall_verdict", ""),
        "audit_status": audit.get("status", ""),
        "feedback_status": feedback.get("status", ""),
        "archive_execution": {
            "status": archive.get("status", ""),
            "confirmed": False,
        },
        "sqlite_location": "runtime/state/project_manager.sqlite3",
        "artifacts": sorted((resumed.get("artifacts") or {}).keys()),
    }


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Run an offline, synthetic local agent judgement dry-run."
    )
    parser.add_argument("--root", required=True)
    parser.add_argument("--use-existing-synthetic-source", action="store_true")
    parser.add_argument("--with-independent-archive-binding", action="store_true")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        result = run_local_agent_dry_run(
            args.root,
            host=DeterministicSyntheticAgentHost(),
            use_existing_synthetic_source=args.use_existing_synthetic_source,
            include_archive_binding=args.with_independent_archive_binding,
        )
    except Exception:
        print(
            json.dumps(
                {"status": "failed", "error_code": "DRY_RUN.UNEXPECTED"},
                sort_keys=True,
            ),
            file=sys.stderr,
        )
        return 1
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0 if result.get("status") == "needs_review" else 2


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = [
    "AgentHost",
    "DeterministicSyntheticAgentHost",
    "DryRunSafetyError",
    "main",
    "run_local_agent_dry_run",
]
