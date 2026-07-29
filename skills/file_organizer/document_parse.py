"""Native-first document parsing for the File Organizer Agent."""

from __future__ import annotations

import hashlib
import json
import mimetypes
import re
from pathlib import Path
from typing import Any, Callable, Mapping

from contracts.structured_document import build_structured_document


NativeParser = Callable[[str], Mapping[str, Any]]
_PROJECT_CODE = re.compile(r"(?:项目(?:编号|代码)?|project(?:\s+code)?)\s*[：:#]?\s*([A-Za-z]{2,12}-\d{1,12})", re.IGNORECASE)
_CONTRACT_CODE = re.compile(r"(?:合同(?:编号|代码)?|contract(?:\s+code)?)\s*[：:#]?\s*([A-Za-z]{1,12}-\d{1,12})", re.IGNORECASE)


class DocumentParseSkill:
    """Read one explicitly supplied file, preferring native extraction to OCR."""

    def __init__(self, native_parser: NativeParser | None = None, ocr_chain: Any = None) -> None:
        self._native_parser = native_parser or parse_native_document
        self._ocr_chain = ocr_chain

    def parse(
        self, source_path: str, *, source_ref: Mapping[str, str] | None = None
    ) -> dict[str, Any]:
        path = Path(source_path)
        if not path.is_file():
            return _review_result(path, source_ref, "INPUT.FILE_UNAVAILABLE")
        resolved_source_ref = _source_ref(path, source_ref)
        content_hash = _sha256(path)
        media_type = _media_type(path)
        try:
            native = dict(self._native_parser(str(path)))
        except Exception:
            native = {"status": "blocked", "reason": "NATIVE.PARSE_FAILED"}

        native_text = native.get("text") if isinstance(native.get("text"), str) else ""
        if native.get("status") == "success" and native_text.strip():
            return build_structured_document(
                source_ref=resolved_source_ref,
                content_hash=content_hash,
                media_type=media_type,
                parser=str(native.get("parser") or "native"),
                text=native_text,
                pages=_list_of_dicts(native.get("pages")),
                tables=_list_of_dicts(native.get("tables")),
                fields=_extract_fields(native_text, native.get("fields")),
            )

        ocr_result = self._extract_ocr(str(path))
        ocr_text = ocr_result.get("text") if isinstance(ocr_result.get("text"), str) else ""
        if ocr_result.get("status") == "success" and ocr_text.strip():
            return build_structured_document(
                source_ref=resolved_source_ref,
                content_hash=content_hash,
                media_type=media_type,
                parser=str(ocr_result.get("provider") or ocr_result.get("engine") or "ocr"),
                text=ocr_text,
                pages=_list_of_dicts(ocr_result.get("pages")),
                tables=[],
                fields=_extract_fields(ocr_text, None),
            )
        return _review_result(
            path,
            resolved_source_ref,
            str(ocr_result.get("reason") or "OCR.PROVIDERS_UNAVAILABLE"),
            content_hash=content_hash,
            media_type=media_type,
        )

    def _extract_ocr(self, source_path: str) -> dict[str, Any]:
        extractor = getattr(self._ocr_chain, "extract", None)
        if not callable(extractor):
            return {"status": "needs_review", "reason": "OCR.PROVIDERS_UNAVAILABLE"}
        try:
            result = extractor(source_path)
        except Exception:
            return {"status": "needs_review", "reason": "OCR.PROVIDERS_UNAVAILABLE"}
        return dict(result) if isinstance(result, Mapping) else {"status": "needs_review", "reason": "OCR.PROVIDERS_UNAVAILABLE"}


def parse_native_document(source_path: str) -> dict[str, Any]:
    """Extract local content without installing engines or calling remote services."""
    path = Path(source_path)
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md", ".csv", ".json", ".xml"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            return {"status": "success", "parser": f"native_{suffix[1:]}", "text": text, "pages": [], "tables": []}
        if suffix == ".docx":
            from docx import Document

            document = Document(path)
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            tables = [
                {"rows": [[cell.text for cell in row.cells] for row in table.rows]}
                for table in document.tables
            ]
            return {"status": "success", "parser": "native_docx", "text": "\n".join(paragraphs), "pages": [], "tables": tables}
        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            tables = []
            chunks = []
            for worksheet in workbook.worksheets:
                rows = [["" if cell is None else str(cell) for cell in row] for row in worksheet.iter_rows(values_only=True)]
                tables.append({"sheet": worksheet.title, "rows": rows})
                chunks.extend("\t".join(row) for row in rows if any(row))
            workbook.close()
            return {"status": "success", "parser": "native_xlsx", "text": "\n".join(chunks), "pages": [], "tables": tables}
        if suffix == ".pdf":
            import fitz

            document = fitz.open(path)
            try:
                pages = [{"page": index, "text": page.get_text("text").strip(), "confidence": None} for index, page in enumerate(document, start=1)]
            finally:
                document.close()
            text = "\n".join(page["text"] for page in pages if page["text"]).strip()
            if text:
                return {"status": "success", "parser": "native_pdf", "text": text, "pages": pages, "tables": []}
            return {"status": "blocked", "reason": "NATIVE.NO_TEXT"}
    except Exception:
        return {"status": "blocked", "reason": "NATIVE.PARSE_FAILED"}
    return {"status": "blocked", "reason": "NATIVE.UNSUPPORTED_MEDIA"}


def _source_ref(path: Path, source_ref: Mapping[str, str] | None) -> dict[str, str]:
    if source_ref is not None:
        return {key: str(source_ref.get(key, "")).strip() for key in ("binding_id", "logical_uri", "storage_provider", "object_key")}
    return {"binding_id": "local", "logical_uri": f"local://{path.name}", "storage_provider": "local", "object_key": path.name}


def _review_result(
    path: Path,
    source_ref: Mapping[str, str] | None,
    reason: str,
    *,
    content_hash: str = "",
    media_type: str = "",
) -> dict[str, Any]:
    return {
        "schema_version": "structured_document.v1",
        "status": "needs_review",
        "reason": reason,
        "source_ref": _source_ref(path, source_ref),
        "content_hash": content_hash or (_sha256(path) if path.is_file() else ""),
        "media_type": media_type or _media_type(path),
        "parser": "unavailable",
        "text": "",
        "pages": [],
        "tables": [],
        "fields": {},
    }


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file_handle:
        for chunk in iter(lambda: file_handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _media_type(path: Path) -> str:
    custom = {".md": "text/markdown", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    return custom.get(path.suffix.lower()) or mimetypes.guess_type(path.name)[0] or "application/octet-stream"


def _list_of_dicts(value: Any) -> list[dict[str, Any]]:
    return [dict(item) for item in value] if isinstance(value, list) and all(isinstance(item, Mapping) for item in value) else []


def _extract_fields(text: str, supplied: Any) -> dict[str, Any]:
    fields = dict(supplied) if isinstance(supplied, Mapping) else {}
    project = _PROJECT_CODE.search(text)
    contract = _CONTRACT_CODE.search(text)
    if project and "project_code" not in fields:
        fields["project_code"] = project.group(1).upper()
    if contract and "contract_code" not in fields:
        fields["contract_code"] = contract.group(1).upper()
    return fields
